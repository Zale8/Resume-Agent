# -*- coding: utf-8 -*-
"""test_content_parser.py — resume.md（内容层）正式解析器测试。

覆盖契约（对应审计报告 P1「建立正式解析器 + 补内容层格式校验」）：
    1. 规范 resume.md → **零诊断**（0 ERROR / 0 WARNING / 0 INFO）；
    2. 内容零丢失：无法归属的行必须出现在 diagnostics 里，绝不静默丢弃；
    3. 致命缺陷（缺姓名 / 空内容 / 文件不存在）→ ResumeParseError；
    4. 章节别名、合并节（专业技能 / 证书奖项）正确归位；
    5. 条目标题四种写法 + 三种异常（缺时间 / 时间非日期 / 完全不匹配）；
    6. 个人信息两种写法（两列表格 / bullet 内联键值）；
    7. 解析产物可直接进入 build_document 端到端渲染。

所有人物 / 学校 / 公司 / 号码均为**虚构夹具**，不含任何真实个人数据。
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from resume_generator.content_parser import (  # noqa: E402
    Diagnostics,
    ParsedResume,
    ResumeParseError,
    SEV_ERROR,
    SEV_INFO,
    SEV_WARNING,
    parse_blocks,
    parse_resume_md,
    parse_resume_md_text,
)
from resume_generator.layout_kit import ResumeBlocks  # noqa: E402


# ---------------------------------------------------------------------------
# 虚构夹具
# ---------------------------------------------------------------------------

CANONICAL = """# 简历内容：李雷 → 演示科技 - 示例岗位

> 生成日期：2026-01-01
> 数据来源：简历库（所有内容均可追溯）

## 个人信息

| 项 | 值 |
|---|---|
| 姓名 | 李雷 |
| 意向 | 示例岗位 ｜ 备用岗位 |
| 电话 | 13800000000 |
| 邮箱 | lilei@example.com |
| 城市 | 示例市　演示市 |
| 出生年月 | 2001.01 |
| 地址 | 示例省示例市 |
| 照片 | 00_个人信息/photos/demo.jpg |

## 个人优势

示例优势第一句。示例优势第二句。

## 教育经历

示例大学｜示例专业（本科）｜2021.09-2025.06｜GPA 3.50/4.0
主修课程：课程甲、课程乙

## 实习经历

### 虚构科技有限公司 - 示例实习生（2024.01-2024.06）
- 完成示例任务甲，覆盖示例场景
- 完成示例任务乙，量化指标 100

### 演示股份有限公司 - 现场实习生（2023.07-2023.09）
- 参与示例产线点检与记录

## 项目经历

### 示例项目 - 负责人（2023.01-2023.12）
- 主导示例项目从 0 到 1

## 校园经历

- 担任示例职务，组织示例活动

## 专业技能

- 掌握示例技能甲与示例技能乙

## 证书奖项

示例证书甲｜示例证书乙
"""


def _personal_only(name_line: bool = True) -> str:
    rows = "| 姓名 | 李雷 |\n" if name_line else ""
    return (
        "# 个人简历\n\n"
        "## 个人信息\n\n"
        "| 项 | 值 |\n|---|---|\n"
        f"{rows}"
        "| 电话 | 13800000000 |\n\n"
        "## 教育经历\n\n"
        "示例大学｜示例专业（本科）｜2021.09-2025.06\n"
    )


# ---------------------------------------------------------------------------
# 测试
# ---------------------------------------------------------------------------

class TestCanonicalParse(unittest.TestCase):
    """验收基准：规范 resume.md 必须零诊断。"""

    def setUp(self):
        self.parsed = parse_resume_md_text(CANONICAL, strict=False)
        self.b = self.parsed.blocks

    def test_01_canonical_produces_zero_diagnostics(self):
        self.assertTrue(
            self.parsed.diagnostics.is_clean,
            f"规范 resume.md 不应产生任何诊断，实际：\n"
            f"{self.parsed.diagnostics.render()}")
        self.assertEqual(self.parsed.diagnostics.counts(),
                         {SEV_ERROR: 0, SEV_WARNING: 0, SEV_INFO: 0})

    def test_02_personal_fields_and_header_lines(self):
        self.assertEqual(self.b.name, "李雷")
        self.assertEqual(self.b.intent, "示例岗位 ｜ 备用岗位")
        self.assertEqual(len(self.b.contact_lines), 2)
        area, contact = self.b.contact_lines
        self.assertIn("求职城市：示例市", area)
        self.assertIn("出生年月：2001.01", area)
        self.assertIn("籍贯：示例省示例市", area)   # 规范用「籍贯」，不写「地址」
        self.assertIn("电话：13800000000", contact)
        self.assertIn("邮箱：lilei@example.com", contact)

    def test_03_body_sections(self):
        self.assertEqual(self.b.summary, "示例优势第一句。示例优势第二句。")
        self.assertEqual(len(self.b.education_lines), 2)
        self.assertIn("示例大学", self.b.education_lines[0])
        self.assertTrue(self.b.education_lines[1].startswith("主修课程："))
        self.assertEqual(self.b.campus, ["担任示例职务，组织示例活动"])
        self.assertEqual(self.b.skills, ["掌握示例技能甲与示例技能乙"])
        self.assertEqual(self.b.certs, ["示例证书甲", "示例证书乙"])

    def test_04_entry_blocks_carry_role_and_meta(self):
        self.assertEqual(len(self.b.internships), 2)
        first = self.b.internships[0]
        self.assertEqual(first.title, "虚构科技有限公司")
        self.assertEqual(first.role, "示例实习生")
        self.assertEqual(first.meta, "2024.01-2024.06")
        self.assertEqual(len(first.bullets), 2)

        self.assertEqual(len(self.b.projects), 1)
        proj = self.b.projects[0]
        self.assertEqual(proj.title, "示例项目")
        self.assertEqual(proj.role, "负责人")
        self.assertEqual(proj.meta, "2023.01-2023.12")

    def test_05_photo_path_kept_when_no_root(self):
        self.assertIsNotNone(self.b.photo_path)
        self.assertTrue(self.b.photo_path.endswith("demo.jpg"))

    def test_06_summary_line_is_readable(self):
        line = self.parsed.summary_line()
        self.assertIn("李雷", line)
        self.assertIn("实习 2", line)


class TestFatal(unittest.TestCase):
    """致命缺陷必须显式报错，不得降级为静默。"""

    def test_10_missing_name_is_fatal(self):
        text = _personal_only(name_line=False)
        with self.assertRaises(ResumeParseError):
            parse_resume_md_text(text)
        parsed = parse_resume_md_text(text, strict=False)
        self.assertTrue(parsed.diagnostics.has_errors)
        self.assertIn("missing_name", parsed.diagnostics.codes())
        self.assertEqual(parsed.blocks.name, "")

    def test_11_h1_stopword_is_not_taken_as_name(self):
        """H1 是「个人简历」这类标题词时，不得被猜成姓名。"""
        parsed = parse_resume_md_text(
            "# 个人简历\n\n## 个人信息\n\n| 项 | 值 |\n|---|---|\n"
            "| 电话 | 13800000000 |\n\n## 教育经历\n\n示例大学\n",
            strict=False)
        self.assertEqual(parsed.blocks.name, "")
        self.assertIn("missing_name", parsed.diagnostics.codes())

    def test_12_empty_body_is_fatal(self):
        text = ("## 个人信息\n\n| 项 | 值 |\n|---|---|\n| 姓名 | 李雷 |\n")
        with self.assertRaises(ResumeParseError):
            parse_resume_md_text(text)
        parsed = parse_resume_md_text(text, strict=False)
        self.assertIn("empty_content", parsed.diagnostics.codes())

    def test_13_missing_file_is_fatal(self):
        with tempfile.TemporaryDirectory() as td:
            missing = Path(td) / "不存在.md"
            with self.assertRaises(ResumeParseError):
                parse_resume_md(missing)
            parsed = parse_resume_md(missing, strict=False)
        self.assertIn("file_not_found", parsed.diagnostics.codes())

    def test_14_raise_if_fatal_returns_self_when_clean(self):
        parsed = parse_resume_md_text(CANONICAL, strict=False)
        self.assertIs(parsed.raise_if_fatal(), parsed)


class TestNoSilentDrop(unittest.TestCase):
    """核心契约：识别不到的内容必须出现在 diagnostics 里。"""

    def test_20_unknown_section_reports_header_and_lines(self):
        text = (
            "## 个人信息\n\n| 项 | 值 |\n|---|---|\n| 姓名 | 李雷 |\n\n"
            "## 补充说明\n\n"
            "- 这一行不该静默消失\n"
            "另一行也不该消失\n\n"
            "## 教育经历\n\n示例大学\n"
        )
        parsed = parse_resume_md_text(text, strict=False)
        codes = parsed.diagnostics.codes()
        self.assertIn("unknown_section", codes)
        self.assertIn("orphan_line", codes)
        raws = " ".join(d.raw for d in parsed.diagnostics.items)
        self.assertIn("这一行不该静默消失", raws)
        self.assertIn("另一行也不该消失", raws)

    def test_21_non_bullet_line_in_entry_is_kept(self):
        text = (
            "## 个人信息\n\n| 项 | 值 |\n|---|---|\n| 姓名 | 李雷 |\n\n"
            "## 项目经历\n\n"
            "### 示例项目 - 负责人（2023.01-2023.12）\n"
            "- 正常要点\n"
            "**项目成果**：指标 100｜指标 200\n"
        )
        parsed = parse_resume_md_text(text, strict=False)
        self.assertIn("non_bullet_line_kept", parsed.diagnostics.codes())
        bullets = parsed.blocks.projects[0].bullets
        self.assertEqual(len(bullets), 2)
        self.assertIn("**项目成果**：指标 100｜指标 200", bullets)

    def test_22_unknown_personal_field_is_reported(self):
        text = (
            "## 个人信息\n\n| 项 | 值 |\n|---|---|\n"
            "| 姓名 | 李雷 |\n| 星座 | 示例座 |\n\n"
            "## 教育经历\n\n示例大学\n"
        )
        parsed = parse_resume_md_text(text, strict=False)
        self.assertIn("unknown_personal_field", parsed.diagnostics.codes())
        raws = " ".join(d.raw for d in parsed.diagnostics.items)
        self.assertIn("星座", raws)

    def test_23_diagnostics_render_lists_codes(self):
        text = (
            "## 个人信息\n\n| 项 | 值 |\n|---|---|\n| 姓名 | 李雷 |\n\n"
            "## 未知章节\n\n- 内容\n"
        )
        parsed = parse_resume_md_text(text, strict=False)
        report = parsed.diagnostics.render()
        self.assertIn("unknown_section", report)
        self.assertIn("WARNING", report)


class TestSectionAliases(unittest.TestCase):
    def test_30_summary_aliases_merge_and_warn(self):
        text = (
            "## 个人信息\n\n| 项 | 值 |\n|---|---|\n| 姓名 | 李雷 |\n\n"
            "## 个人简介\n\n简介内容。\n\n"
            "## 自我评价\n\n评价内容。\n"
        )
        parsed = parse_resume_md_text(text, strict=False)
        self.assertIn("简介内容", parsed.blocks.summary)
        self.assertIn("评价内容", parsed.blocks.summary)
        self.assertIn("alias_section_merge", parsed.diagnostics.codes())

    def test_31_certs_aliases(self):
        for title in ("证书奖项", "荣誉奖项", "获奖情况"):
            text = (
                "## 个人信息\n\n| 项 | 值 |\n|---|---|\n| 姓名 | 李雷 |\n\n"
                f"## {title}\n\n- 示例证书\n"
            )
            parsed = parse_resume_md_text(text, strict=False)
            self.assertEqual(parsed.blocks.certs, ["示例证书"], title)
            self.assertTrue(parsed.diagnostics.is_clean,
                            f"{title} 应为已知别名：{parsed.diagnostics.render()}")

    def test_32_pending_section_is_registered_not_warned(self):
        text = (
            "## 个人信息\n\n| 项 | 值 |\n|---|---|\n| 姓名 | 李雷 |\n\n"
            "## 教育经历\n\n示例大学\n\n"
            "## 待补充\n\n- 示例 Gap 项\n"
        )
        parsed = parse_resume_md_text(text, strict=False)
        self.assertFalse(parsed.diagnostics.warnings)
        self.assertIn("section_not_rendered", parsed.diagnostics.codes())

    def test_33_merged_skills_certs_section(self):
        text = (
            "## 个人信息\n\n| 项 | 值 |\n|---|---|\n| 姓名 | 李雷 |\n\n"
            "## 专业技能 / 证书奖项\n\n"
            "- 技能：技能甲\n"
            "- 证书：证书甲｜证书乙\n"
        )
        parsed = parse_resume_md_text(text, strict=False)
        self.assertEqual(parsed.blocks.skills, ["技能甲"])
        self.assertEqual(parsed.blocks.certs, ["证书甲", "证书乙"])

    def test_34_personal_inline_bullets(self):
        text = (
            "## 基本信息\n\n"
            "- 姓名：李雷\n"
            "- 出生年月：2001.01 | 民族：示例族\n\n"
            "## 教育经历\n\n示例大学\n"
        )
        parsed = parse_resume_md_text(text, strict=False)
        b = parsed.blocks
        self.assertEqual(b.name, "李雷")
        area = b.contact_lines[0]
        self.assertIn("出生年月：2001.01", area)
        # 民族已识别但无槽位 → INFO 登记，且在 raw 里留下痕迹
        self.assertIn("unmapped_personal_field", parsed.diagnostics.codes())
        raws = " ".join(d.raw for d in parsed.diagnostics.items)
        self.assertIn("民族", raws)


class TestEntryTitles(unittest.TestCase):
    def _parse_entry(self, title: str) -> ParsedResume:
        text = (
            "## 个人信息\n\n| 项 | 值 |\n|---|---|\n| 姓名 | 李雷 |\n\n"
            "## 实习经历\n\n"
            f"### {title}\n"
            "- 要点\n"
        )
        return parse_resume_md_text(text, strict=False)

    def test_40_dash_with_fullwidth_paren(self):
        p = self._parse_entry("甲公司 - 乙岗位（2020.01-2021.02）")
        item = p.blocks.internships[0]
        self.assertEqual((item.title, item.role, item.meta),
                         ("甲公司", "乙岗位", "2020.01-2021.02"))
        self.assertTrue(p.diagnostics.is_clean, p.diagnostics.render())

    def test_41_pipe_three_parts_with_spaces_in_date(self):
        p = self._parse_entry("丙公司 | 丁岗位 | 2022.01 - 2023.02")
        item = p.blocks.internships[0]
        self.assertEqual((item.title, item.role, item.meta),
                         ("丙公司", "丁岗位", "2022.01 - 2023.02"))
        self.assertTrue(p.diagnostics.is_clean, p.diagnostics.render())

    def test_42_fullwidth_pipe_two_parts_has_no_meta(self):
        p = self._parse_entry("戊项目｜负责人")
        item = p.blocks.internships[0]
        self.assertEqual((item.title, item.role, item.meta), ("戊项目", "负责人", ""))
        self.assertIn("entry_without_meta", p.diagnostics.codes())

    def test_43_unparsable_title_is_warned(self):
        p = self._parse_entry("一个完全不符合格式的标题")
        item = p.blocks.internships[0]
        self.assertEqual(item.title, "一个完全不符合格式的标题")
        self.assertIn("unparsed_entry_title", p.diagnostics.codes())

    def test_44_meta_that_is_not_a_date_is_warned(self):
        p = self._parse_entry("己公司 - 庚岗位（去年夏天）")
        self.assertIn("entry_meta_not_date", p.diagnostics.codes())

    def test_45_to_present_meta_is_accepted(self):
        p = self._parse_entry("辛公司 - 壬岗位（2026.07-至今）")
        self.assertTrue(p.diagnostics.is_clean, p.diagnostics.render())


class TestDiagnosticsApi(unittest.TestCase):
    def test_50_counts_and_codes(self):
        d = Diagnostics()
        d.error("e1", "错误")
        d.warn("w1", "警告")
        d.info("i1", "提示")
        self.assertEqual(d.counts(), {SEV_ERROR: 1, SEV_WARNING: 1, SEV_INFO: 1})
        self.assertTrue(d.has_errors)
        self.assertFalse(d.is_clean)
        self.assertEqual(d.codes(SEV_INFO), ["i1"])
        self.assertEqual(d.codes(), ["e1", "w1", "i1"])

    def test_51_render_groups_by_code(self):
        d = Diagnostics()
        for i in range(3):
            d.warn("same", f"第 {i} 条")
        report = d.render()
        self.assertIn("same × 3", report)
        self.assertNotIn("ERROR 1", report)

    def test_52_empty_render(self):
        self.assertIn("无", Diagnostics().render())


class TestPhotoResolution(unittest.TestCase):
    def test_60_existing_photo_is_resolved_without_warning(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            (root / "photos").mkdir()
            (root / "photos" / "demo.jpg").write_bytes(b"x")
            text = (
                "## 个人信息\n\n| 项 | 值 |\n|---|---|\n| 姓名 | 李雷 |\n"
                "| 照片 | photos/demo.jpg |\n\n"
                "## 教育经历\n\n示例大学\n"
            )
            parsed = parse_resume_md_text(text, photo_root=root, strict=False)
            self.assertEqual(
                Path(parsed.blocks.photo_path),
                root / "photos" / "demo.jpg")
            self.assertNotIn("photo_not_found", parsed.diagnostics.codes())

    def test_61_missing_photo_is_warned(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            text = (
                "## 个人信息\n\n| 项 | 值 |\n|---|---|\n| 姓名 | 李雷 |\n"
                "| 照片 | photos/nope.jpg |\n\n"
                "## 教育经历\n\n示例大学\n"
            )
            parsed = parse_resume_md_text(text, photo_root=root, strict=False)
            self.assertIn("photo_not_found", parsed.diagnostics.codes())


class TestHelpers(unittest.TestCase):
    def test_70_parse_blocks_from_text_and_path(self):
        blocks = parse_blocks(CANONICAL)
        self.assertIsInstance(blocks, ResumeBlocks)
        self.assertEqual(blocks.name, "李雷")

        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "resume.md"
            p.write_text(CANONICAL, encoding="utf-8")
            from_file = parse_blocks(p)
            self.assertEqual(from_file.name, blocks.name)
            self.assertEqual(len(from_file.internships), len(blocks.internships))

    def test_71_parse_resume_md_from_file(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "resume.md"
            p.write_text(CANONICAL, encoding="utf-8")
            parsed = parse_resume_md(p)
            self.assertEqual(parsed.source, p)
            self.assertEqual(parsed.blocks.name, "李雷")


class TestEndToEndRender(unittest.TestCase):
    """解析产物可直接进入渲染器（无 python-docx 时跳过）。"""

    def test_80_parsed_blocks_render_to_docx(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("未安装 python-docx")

        from resume_generator.design.presets import build_p3_spec
        from resume_generator.skeletons import build_document

        parsed = parse_resume_md_text(CANONICAL, strict=False)
        doc = build_document(parsed.blocks, design_spec=build_p3_spec())
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "out.docx"
            doc.save(str(out))
            self.assertTrue(out.is_file())
            self.assertGreater(out.stat().st_size, 0)

            # 姓名与优势必须真的落到文档里（不丢内容）
            texts = [p.text for p in doc.paragraphs]
            joined = "\n".join(texts)
            for para in doc.tables:
                for row in para.rows:
                    for cell in row.cells:
                        joined += "\n" + "\n".join(p.text for p in cell.paragraphs)
            self.assertIn("李雷", joined)
            self.assertIn("示例优势第一句", joined)


if __name__ == "__main__":
    unittest.main()
