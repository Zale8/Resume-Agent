# -*- coding: utf-8 -*-
"""test_renderer_palette.py — Phase 3B-2 Palette / Color 接线回归。

验证：
    T1 旧路径三骨架颜色与改造前逐值一致（深块三固定色 / 浅侧栏 / 发丝线）。
    T2 P1 ColorSpec 的 8 个角色真实进入 Renderer 并落到 DOCX XML。
    T3 P2 ColorSpec（浅侧栏范式）正确进入 Renderer；sidebar_heading 走
       Profile fallback（Schema Gap 登记）。
    T4 Spec 显式缺色（None）→ 骨架 Profile fallback，登记 <not_defined>。
    T5 Spec undetermined → Profile fallback，登记 colors.<role>，不编造；
       并验证 fallback 色真实落到 XML。
    T6 真实 DOCX XML（run 色 / shd 填充 / pBdr 边框）与 RenderStyle 一致。
    T7 源码守卫：skeletons.py 代码 token 中不得再出现裸 hex / RGBColor /
       palette.ink|accent|muted|header|rule 直引。
    T8 Phase 3B-1 Typography 回归套件继续全绿（子进程调用）。

简历数据全部虚构（张三）；不向用户数据目录写任何文件。
用法：python tests/test_renderer_palette.py
"""
from __future__ import annotations

import io
import subprocess
import sys
import tokenize
import unittest
from dataclasses import replace
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from docx.oxml.ns import qn  # noqa: E402

from resume_generator.design.spec import Sourced  # noqa: E402
from resume_generator.layout_kit import (  # noqa: E402
    PALETTES,
    BulletBlock,
    ResumeBlocks,
)
from resume_generator.render_style import (  # noqa: E402
    SKELETON_DEFAULT_STYLE,
    resolve_style_for_spec,
)
from resume_generator.skeletons import (  # noqa: E402
    SKELETON_BANNER,
    SKELETON_MINIMAL,
    SKELETON_SIDEBAR,
    build_document,
    resolve_palette,
)


# ---------------------------------------------------------------------------
# 虚构夹具（与 test_renderer_typography 同口径，关键字构造 BulletBlock）
# ---------------------------------------------------------------------------

def fake_blocks():
    return ResumeBlocks(
        name="张三",
        intent="测试岗位",
        contact_lines=["138-0000-0000", "zhangsan@example.com"],
        summary="虚构的个人优势简介，用于调色板接线测试。",
        education_lines=["XX大学  测试专业（本科）    2023.09 - 2027.06"],
        internships=[
            BulletBlock(
                title="某某有限公司",
                role="测试实习生",
                meta="2025.06 - 2025.09",
                tags="标签甲 · 标签乙",
                bullets=["完成了一项虚构的测试工作。"],
            )
        ],
        projects=[],
        campus=["虚构校园经历甲"],
        skills=["技能甲", "技能乙"],
        certs=["虚构证书甲"],
    )


# ---------------------------------------------------------------------------
# DOCX XML 读取 helper（只观察，不修改）
# ---------------------------------------------------------------------------

def all_paragraphs(doc):
    ps = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                ps.extend(cell.paragraphs)
    return ps


def find_para(doc, text):
    for p in all_paragraphs(doc):
        if text in p.text:
            return p
    raise AssertionError(f"未找到包含 {text!r} 的段落")


def run_color(run):
    rPr = run._r.find(qn("w:rPr"))
    if rPr is None:
        return None
    color = rPr.find(qn("w:color"))
    return None if color is None else color.get(qn("w:val"))


def run_color_of_para(doc, text):
    return run_color(next(r for r in find_para(doc, text).runs if text in r.text))


def border_of_para(p):
    """返回 (color, sz)；无边框返回 (None, None)。"""
    pPr = p._p.find(qn("w:pPr"))
    if pPr is None:
        return None, None
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        return None, None
    bottom = pBdr.find(qn("w:bottom"))
    if bottom is None:
        return None, None
    return bottom.get(qn("w:color")), bottom.get(qn("w:sz"))


def cell_fill(cell):
    tcPr = cell._tc.find(qn("w:tcPr"))
    shd = tcPr.find(qn("w:shd"))
    return None if shd is None else shd.get(qn("w:fill"))


def all_bottom_border_colors(doc):
    return {border_of_para(p)[0]
            for p in all_paragraphs(doc)
            if border_of_para(p)[0] is not None}


def resolve_spec_style(spec, sk):
    """与 _build_from_spec 同口径地解析 RenderStyle（含 RenderPalette）。"""
    return resolve_style_for_spec(
        spec, SKELETON_DEFAULT_STYLE[sk], sk, resolve_palette(None, sk))


# ===========================================================================
# 测试
# ===========================================================================

class TestRendererPalette(unittest.TestCase):

    # -- Test 1：旧路径颜色保值 ---------------------------------------------
    def test_01_legacy_paths_colors_unchanged(self):
        pharma = PALETTES["pharma_teal"]
        steel = PALETTES["manufacturing_steel"]
        minimal = PALETTES["minimal_ink"]

        # banner：深块底 = 色板 header；深块三文字色为骨架固定值
        d = build_document(fake_blocks(), SKELETON_BANNER)
        left, right = d.tables[0].rows[0].cells
        self.assertEqual(cell_fill(left), pharma.header.lstrip("#"))
        self.assertEqual(cell_fill(right), pharma.header.lstrip("#"))
        self.assertEqual(run_color_of_para(d, "张三"), "FFFFFF")
        self.assertEqual(run_color_of_para(d, "测试岗位"), "F7DC6F")
        self.assertEqual(run_color_of_para(d, "138-0000-0000"), "EAECEE")
        # banner 正文/发丝线仍取行业色板
        borders = all_bottom_border_colors(d)
        self.assertIn(pharma.accent.lstrip("#"), borders)   # 强调色条
        self.assertIn(pharma.rule.lstrip("#"), borders)     # 标题发丝线

        # sidebar：左栏浅底 = 色板 rule；姓名深色 = 色板 header
        d = build_document(fake_blocks(), SKELETON_SIDEBAR)
        left, right = d.tables[0].rows[0].cells
        self.assertEqual(cell_fill(left), steel.rule.lstrip("#"))
        self.assertIsNone(cell_fill(right))
        self.assertEqual(run_color_of_para(d, "张三"),
                         steel.header.lstrip("#"))

        # minimal：正文 ink / 顶部强调线 accent / 编号标题下发丝线 rule
        d = build_document(fake_blocks(), SKELETON_MINIMAL)
        self.assertEqual(run_color_of_para(d, "张三"),
                         minimal.ink.lstrip("#"))
        borders = all_bottom_border_colors(d)
        self.assertIn(minimal.accent.lstrip("#"), borders)
        self.assertIn(minimal.rule.lstrip("#"), borders)

    # -- Test 2：P1 Spec Palette 进入 Renderer 并落 XML ---------------------
    def test_02_p1_spec_colors_reach_renderer_and_xml(self):
        from resume_generator.design import build_p1_spec
        spec = build_p1_spec()
        st = resolve_spec_style(spec, SKELETON_BANNER)
        rp = st.palette
        # RenderStyle 层：8 个被消费角色全部取 Spec 值（带 # 原样保留）
        self.assertEqual(rp.ink, "#1A1A1A")
        self.assertEqual(rp.accent, "#E31937")
        self.assertEqual(rp.muted, "#8A8A8A")
        self.assertEqual(rp.dark_block, "#252525")
        self.assertEqual(rp.hairline, "#DFDFDF")
        self.assertEqual(rp.on_dark_primary, "#FFFFFF")
        self.assertEqual(rp.on_dark_secondary, "#E8E8E8")
        self.assertEqual(rp.on_dark_tertiary, "#BBBBBB")
        # P1 不使用浅侧栏；paper 透传备用
        self.assertIsNone(rp.light_sidebar)
        self.assertEqual(rp.paper, "#FFFFFF")

        # 真实 XML 层
        d = build_document(fake_blocks(), design_spec=spec)
        left, right = d.tables[0].rows[0].cells
        self.assertEqual(cell_fill(left), "252525")       # 深块底
        self.assertEqual(cell_fill(right), "252525")
        self.assertEqual(run_color_of_para(d, "张三"), "FFFFFF")
        self.assertEqual(run_color_of_para(d, "测试岗位"), "BBBBBB")
        self.assertEqual(run_color_of_para(d, "138-0000-0000"), "E8E8E8")
        # 强调色条（颜色 = P1 accent；厚度自 3B-3 起由 Spacing Spec 决定，
        # 不再用固定 sz=18 定位）
        bar = next(p for p in d.paragraphs
                   if border_of_para(p)[0] == "E31937")
        self.assertEqual(border_of_para(bar)[0], "E31937")
        # P-5：P1 divider.weight=0pt=禁用 → 板块发丝线不写入 DOCX
        # （RenderPalette.hairline 仍解析为 #DFDFDF，但无渲染消费点）
        self.assertNotIn("DFDFDF", all_bottom_border_colors(d))
        # ink 落到教育行；muted 落到右对齐日期 run
        self.assertEqual(run_color_of_para(d, "XX大学"), "1A1A1A")
        date_p = find_para(d, "2025.06 - 2025.09")
        date_run = next(r for r in date_p.runs if "2025" in r.text)
        self.assertEqual(run_color(date_run), "8A8A8A")

    # -- Test 3：P2 Spec Palette 进入 Renderer ------------------------------
    def test_03_p2_spec_colors_reach_renderer_and_xml(self):
        from resume_generator.design import build_p2_spec
        spec = build_p2_spec()
        st = resolve_spec_style(spec, SKELETON_SIDEBAR)
        rp = st.palette
        self.assertEqual(rp.ink, "#222222")
        self.assertEqual(rp.accent, "#24406B")
        self.assertEqual(rp.muted, "#888888")
        self.assertEqual(rp.light_sidebar, "#EFF2F7")
        self.assertEqual(rp.hairline, "#DDE2E8")
        self.assertEqual(rp.paper, "#FFFFFF")
        # P2 范式不使用深块，三个 on_dark 角色保持 Profile（None）
        self.assertIsNone(rp.dark_block)
        self.assertIsNone(rp.on_dark_primary)
        # sidebar_heading 无 Spec 角色：走骨架默认色板 header（manufacturing）
        steel = PALETTES["manufacturing_steel"]
        self.assertEqual(rp.sidebar_heading, steel.header)
        # Schema Gap 必须登记
        self.assertIn("colors.sidebar_heading<schema_gap>",
                      st.spec_fallbacks)

        d = build_document(fake_blocks(), design_spec=spec)
        left, right = d.tables[0].rows[0].cells
        self.assertEqual(cell_fill(left), "EFF2F7")        # 浅侧栏底
        self.assertEqual(run_color_of_para(d, "测试岗位"), "24406B")
        self.assertEqual(run_color_of_para(d, "张三"),
                         steel.header.lstrip("#"))         # schema gap fallback
        # 侧栏分组标签下划线 = accent；主栏标题下划线 = hairline
        side_borders = {border_of_para(p)[0]
                        for p in left.paragraphs
                        if border_of_para(p)[0]}
        main_borders = {border_of_para(p)[0]
                        for p in right.paragraphs
                        if border_of_para(p)[0]}
        self.assertIn("24406B", side_borders)
        self.assertIn("DDE2E8", main_borders)
        # 主栏教育行 ink；日期 muted
        self.assertEqual(run_color_of_para(d, "XX大学"), "222222")
        date_p = find_para(d, "2025.06 - 2025.09")
        date_run = next(r for r in date_p.runs if "2025" in r.text)
        self.assertEqual(run_color(date_run), "888888")

    # -- Test 4：Spec 缺色（None）→ fallback + not_defined 登记 -------------
    def test_04_missing_color_falls_back_with_log(self):
        from resume_generator.design import build_p1_spec
        spec = build_p1_spec()
        spec = replace(spec, colors=replace(spec.colors, muted=None))
        st = resolve_spec_style(spec, SKELETON_BANNER)
        pharma = PALETTES["pharma_teal"]
        # muted 未定义 → 回退骨架默认行业色板（非 Spec 值 8A8A8A）
        self.assertEqual(st.palette.muted, pharma.muted)
        self.assertIn("colors.muted<not_defined>", st.spec_fallbacks)
        # 其它已定义角色不受影响
        self.assertEqual(st.palette.accent, "#E31937")

    # -- Test 5：undetermined → fallback + 登记，且真实落到 XML -------------
    def test_05_undetermined_color_falls_back_with_log(self):
        from resume_generator.design import build_p1_spec
        spec = build_p1_spec()
        spec = replace(
            spec,
            colors=replace(spec.colors, accent=Sourced.undetermined(
                "测试用：强调色知识缺口")))
        st = resolve_spec_style(spec, SKELETON_BANNER)
        pharma = PALETTES["pharma_teal"]
        self.assertEqual(st.palette.accent, pharma.accent)
        # undetermined 登记不带 <not_defined> 后缀（与 Typography 口径一致）
        self.assertIn("colors.accent", st.spec_fallbacks)
        self.assertNotIn("colors.accent<not_defined>", st.spec_fallbacks)

        # Validator 只给 WARNING：完整 build 仍成功，fallback 色落到色条
        # （按颜色定位；色条厚度自 3B-3 起不再固定为 sz=18）
        d = build_document(fake_blocks(), design_spec=spec)
        bar = next(p for p in d.paragraphs
                   if border_of_para(p)[0]
                   == pharma.accent.lstrip("#"))
        self.assertEqual(border_of_para(bar)[0],
                         pharma.accent.lstrip("#"))

    # -- Test 6：XML 与 RenderStyle 三类元素一致性 --------------------------
    def test_06_xml_matches_renderstyle(self):
        from resume_generator.design import build_p1_spec
        spec = build_p1_spec()
        st = resolve_spec_style(spec, SKELETON_BANNER)
        d = build_document(fake_blocks(), design_spec=spec)

        # 字体颜色（深块姓名）
        self.assertEqual(run_color_of_para(d, "张三"),
                         st.palette.on_dark_primary.lstrip("#"))
        # 背景填充
        left = d.tables[0].rows[0].cells[0]
        self.assertEqual(cell_fill(left),
                         st.palette.dark_block.lstrip("#"))
        # 边框（强调色条按 accent 颜色定位；厚度归 3B-3 Spacing 测试断言）
        bar = next(p for p in d.paragraphs
                   if border_of_para(p)[0]
                   == st.palette.accent.lstrip("#"))
        self.assertEqual(border_of_para(bar)[0],
                         st.palette.accent.lstrip("#"))
        # P-5：divider 0pt 禁用后 hairline 色不落 XML（色仍在 RenderPalette）
        self.assertNotIn(st.palette.hairline.lstrip("#"),
                         all_bottom_border_colors(d))

    # -- Test 7：源码守卫 ---------------------------------------------------
    def test_07_no_hardcoded_color_in_skeletons(self):
        src_path = SRC / "resume_generator" / "skeletons.py"
        code_lines = self._code_lines_without_strings_comments(src_path)
        bad_hex = []
        bad_rgb = []
        bad_direct = []
        for lineno, text in code_lines.items():
            import re
            if re.search(r"#[0-9A-Fa-f]{6}\b", text):
                bad_hex.append((lineno, text.strip()))
            if "RGBColor" in text:
                bad_rgb.append((lineno, text.strip()))
            # Renderer 只能读 RenderPalette（st.palette / style.palette），
            # 不得再直接读行业 Palette 的语义字段
            if re.search(r"(?<![.\w])palette\.(ink|accent|muted|paper|header|rule)\b",
                         text):
                bad_direct.append((lineno, text.strip()))
        self.assertEqual(bad_hex, [], f"skeletons.py 出现裸 hex：{bad_hex}")
        self.assertEqual(bad_rgb, [], f"skeletons.py 出现 RGBColor：{bad_rgb}")
        self.assertEqual(bad_direct, [],
                         f"skeletons.py 直引 Palette 语义字段：{bad_direct}")

    @staticmethod
    def _code_lines_without_strings_comments(path: Path):
        """用 tokenize 剥除 STRING/COMMENT，返回 {行号: 该行代码 token 拼接}。

        docstring 内的历史颜色描述不算硬编码，因此必须按 token 剥离，
        不能直接对源文本做正则。
        """
        lines: dict = {}
        with tokenize.open(path) as f:
            tokens = tokenize.generate_tokens(
                io.StringIO(f.read()).readline)
            for tok in tokens:
                if tok.type in (tokenize.STRING, tokenize.COMMENT,
                                tokenize.ENCODING):
                    continue
                lines.setdefault(tok.start[0], []).append(tok.string)
        return {n: " ".join(parts) for n, parts in lines.items()}

    # -- Test 8：Typography 回归（3B-1 套件必须继续全绿） -------------------
    def test_08_typography_regression_suite(self):
        proc = subprocess.run(
            [sys.executable, str(ROOT / "tests" /
                                 "test_renderer_typography.py")],
            cwd=str(ROOT),
            capture_output=True,
            text=True,
        )
        self.assertEqual(
            proc.returncode, 0,
            msg=f"typography 回归失败：\n{proc.stdout}\n{proc.stderr}")


if __name__ == "__main__":
    unittest.main(verbosity=2)
