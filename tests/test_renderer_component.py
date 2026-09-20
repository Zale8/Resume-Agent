# -*- coding: utf-8 -*-
"""test_renderer_component.py — Phase 3B-4 Component / 组件结构参数回归。

验证：
    解析（Resolver）
      1. Spec 明确值 → RenderComponents（角色分隔符 / bullet 符号 / 标题符号 /
         日期模式 / 身份对齐 / 等比保护）
      2. Spec 缺失（None）→ 骨架 Profile fallback，登记 <not_defined>
      3. Spec undetermined → 骨架 Profile fallback，登记裸键名
      4. Schema Gap（Spec 无此设计概念）→ 永久 Profile，登记 <schema_gap>
      5. 三骨架默认 Profile = 改造前逐字符审计值，旧路径挂载 components
    渲染（真实 DOCX）
      6. SectionTitle：▍符号式 / 01 编号式 / Spec 自定义符号
      7. ExperienceHeader：分隔符、日期 \\t 前缀与右制表位、minimal 双空格
      8. Bullet：旧路径 "· "，Spec 路径 experience.bullet.symbol="•"
      9. Tag：当前仅为纯文本 muted 行（无 border/fill）— 锁定现状能力
      10. Photo：照片插入 + preserve_aspect 槽位；缺图不插不抛
      11. AccentBar：空段落 + 底边框，颜色来自 Palette / 厚度来自 Spacing
      12. ContactBlock：普通文本段落，无图标 / 无边框，left 时不写 w:jc
    另含源码守卫：组件结构字面量不得再散落 skeletons.py。

简历数据全部虚构（张三）；Document 默认只在内存构建，照片用 1px 临时 PNG。
用法：python tests/test_renderer_component.py
"""
from __future__ import annotations

import io
import subprocess
import sys
import tempfile
import tokenize
import unittest
from dataclasses import replace
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from docx.oxml.ns import qn  # noqa: E402

from resume_generator.design.spec import Sourced  # noqa: E402
from resume_generator.layout_kit import (  # noqa: E402
    BulletBlock,
    ResumeBlocks,
)
from resume_generator.render_style import (  # noqa: E402
    SKELETON_DEFAULT_STYLE,
    resolve_style_for_spec,
)
from resume_generator.skeletons import (  # noqa: E402
    SKELETON_BANNER,
    SKELETON_MINIMAL,
    SKELETON_SIDEBAR,
    build_document,
    resolve_palette,
)


# ---------------------------------------------------------------------------
# 虚构夹具（与 typography/spacing 套件同口径，BulletBlock 一律关键字构造）
# ---------------------------------------------------------------------------

def fake_blocks(photo_path=None):
    return ResumeBlocks(
        name="张三",
        intent="测试岗位",
        contact_lines=["138-0000-0000", "zhangsan@example.com"],
        summary="虚构的个人优势简介，用于组件接线测试。",
        education_lines=["XX大学  测试专业（本科）    2023.09 - 2027.06"],
        internships=[
            BulletBlock(
                title="某某有限公司",
                role="测试实习生",
                meta="2025.06 - 2025.09",
                tags="标签甲 · 标签乙",
                bullets=["完成了一项虚构的测试工作。"],
            )
        ],
        projects=[],
        campus=["虚构校园经历甲"],
        skills=["技能甲", "技能乙"],
        certs=["虚构证书甲", "虚构证书乙"],
        photo_path=photo_path,
    )


def resolve_spec_style(spec, sk):
    """与 _build_from_spec 同口径地解析 RenderStyle（含 components）。"""
    return resolve_style_for_spec(
        spec, SKELETON_DEFAULT_STYLE[sk], sk, resolve_palette(None, sk))


def all_paragraphs(doc):
    """body 段落 + 全部表格单元格段落（sidebar 内容在格内）。"""
    out = list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                out.extend(cell.paragraphs)
    return out


def find_para(doc, fragment):
    for p in all_paragraphs(doc):
        if fragment in p.text:
            return p
    raise AssertionError(f"未找到包含 {fragment!r} 的段落")


def jc_of_para(p):
    pPr = p._p.find(qn("w:pPr"))
    if pPr is None:
        return None
    jc = pPr.find(qn("w:jc"))
    return None if jc is None else jc.get(qn("w:val"))


def has_tab_stop(p):
    return p._p.find(qn("w:pPr")) is not None and \
        p._p.find(qn("w:pPr")).find(qn("w:tabs")) is not None


def bottom_border(p):
    pPr = p._p.find(qn("w:pPr"))
    if pPr is None:
        return None
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        return None
    bottom = pBdr.find(qn("w:bottom"))
    if bottom is None:
        return None
    return {"sz": bottom.get(qn("w:sz")),
            "color": bottom.get(qn("w:color")),
            "val": bottom.get(qn("w:val"))}


def has_drawing(p):
    return p._p.find(".//" + qn("w:drawing")) is not None


def run_color(run):
    """run 的 w:color 值（hex 小写无 #）；未写颜色返回 None。"""
    rPr = run._r.find(qn("w:rPr"))
    if rPr is None:
        return None
    color = rPr.find(qn("w:color"))
    return None if color is None else color.get(qn("w:val"))


# 1×1 像素 PNG 动态生成（仅用于让 insert_photo 走真实插入路径，零二进制夹具）
import binascii  # noqa: E402
import struct  # noqa: E402
import zlib  # noqa: E402


def _tiny_png_bytes():
    def chunk(tag, data):
        crc = binascii.crc32(tag + data) & 0xFFFFFFFF
        return struct.pack(">I", len(data)) + tag + data + struct.pack(
            ">I", crc)

    raw = b"\x00\xff\x00\x00"  # 1 行：filter byte + RGB
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b""))


_TINY_PNG = _tiny_png_bytes()


# ===========================================================================
# 一、Resolver 解析
# ===========================================================================

class TestComponentResolution(unittest.TestCase):

    # -- Test 1：Spec 明确值进入 RenderComponents ---------------------------
    def test_01_explicit_component_values(self):
        from resume_generator.design import build_p1_spec, build_p2_spec

        st = resolve_spec_style(build_p1_spec(), SKELETON_BANNER)
        cp = st.components
        self.assertEqual(cp.role_separator, "｜")        # P1 全角竖线
        self.assertEqual(cp.bullet_symbol, "•")          # BulletStyle.symbol
        self.assertEqual(cp.section_title_symbol, "▍")   # section.prefix
        self.assertEqual(cp.date_mode, "tab_stop")
        # 身份对齐复合映射
        self.assertEqual(cp.identity_text_align, "left")
        self.assertEqual(cp.photo_align, "right")
        # 等比硬规则（distortion_allowed 必须为 False）
        self.assertIs(cp.photo_preserve_aspect, True)

        st2 = resolve_spec_style(build_p2_spec(), SKELETON_SIDEBAR)
        cp2 = st2.components
        self.assertEqual(cp2.date_mode, "column")
        self.assertEqual(cp2.identity_text_align, "center")
        self.assertEqual(cp2.photo_align, "center")
        self.assertEqual(cp2.bullet_symbol, "•")
        # P2 role.separator 显式 None = 未定义 → Profile，不是空串
        self.assertEqual(cp2.role_separator, "  ·  ")
        self.assertIn("components.role_separator<not_defined>",
                      st2.spec_fallbacks)

    # -- Test 2：Spec 缺失（None）→ Profile + not_defined -------------------
    def test_02_missing_component_falls_back(self):
        from resume_generator.design import build_p1_spec
        spec = build_p1_spec()
        spec = replace(
            spec,
            section=replace(spec.section, prefix=None),
            experience=replace(
                spec.experience,
                role=replace(spec.experience.role, separator=None),
                bullet=replace(spec.experience.bullet, symbol="")))
        st = resolve_spec_style(spec, SKELETON_BANNER)
        cp = st.components
        self.assertEqual(cp.section_title_symbol, "▍")       # Profile
        self.assertEqual(cp.role_separator, "  ·  ")          # Profile
        self.assertEqual(cp.bullet_symbol, "·")               # Profile
        for key in ("components.section_title_symbol<not_defined>",
                    "components.role_separator<not_defined>",
                    "components.bullet_symbol<not_defined>"):
            self.assertIn(key, st.spec_fallbacks)

    # -- Test 3：Spec undetermined → Profile + 裸键登记 ---------------------
    def test_03_undetermined_component_falls_back(self):
        from resume_generator.design import build_p1_spec
        spec = build_p1_spec()
        spec = replace(
            spec,
            section=replace(
                spec.section,
                prefix=Sourced.undetermined("测试用：标题符号知识缺口")))
        st = resolve_spec_style(spec, SKELETON_BANNER)
        self.assertEqual(st.components.section_title_symbol, "▍")
        self.assertIn("components.section_title_symbol", st.spec_fallbacks)
        self.assertNotIn(
            "components.section_title_symbol<not_defined>",
            st.spec_fallbacks)

    # -- Test 4：Schema Gap → 永久 Profile + schema_gap 登记 ----------------
    def test_04_schema_gap_slots_keep_profile(self):
        from resume_generator.design import build_p1_spec, build_p2_spec

        fb = resolve_spec_style(build_p1_spec(),
                                SKELETON_BANNER).spec_fallbacks
        cp = resolve_spec_style(build_p1_spec(),
                                SKELETON_BANNER).components
        # Spec 完全没有这些设计概念，Profile 逐字符保值
        self.assertEqual(cp.date_prefix, "\t")
        self.assertEqual(cp.bullet_gap, " ")
        self.assertEqual(cp.section_title_gap, " ")
        self.assertEqual(cp.cert_label, "证书  ")
        self.assertEqual(cp.cert_separator, "  ·  ")
        for key in ("components.date_prefix<schema_gap>",
                    "components.bullet_gap<schema_gap>",
                    "components.section_title_gap<schema_gap>",
                    "components.cert_label<schema_gap>",
                    "components.cert_separator<schema_gap>"):
            self.assertIn(key, fb)

        # minimal 的编号格式；sidebar 的标签大写变换同为 Schema Gap
        cpm = resolve_spec_style(build_p2_spec(), SKELETON_MINIMAL).components
        self.assertEqual(cpm.title_numbering_template, "%02d")
        self.assertEqual(cpm.title_numbering_gap, "  ")
        self.assertIn("components.title_numbering_template<schema_gap>",
                      resolve_spec_style(build_p2_spec(),
                                         SKELETON_MINIMAL).spec_fallbacks)
        cps = resolve_spec_style(build_p2_spec(), SKELETON_SIDEBAR).components
        self.assertEqual(cps.side_label_transform, "upper_ascii")
        self.assertIn("components.side_label_transform<schema_gap>",
                      resolve_spec_style(build_p2_spec(),
                                         SKELETON_SIDEBAR).spec_fallbacks)

    # -- Test 5：三骨架 Profile 审计值 + 旧路径挂载 -------------------------
    def test_05_skeleton_profiles_audited_values(self):
        b = SKELETON_DEFAULT_STYLE[SKELETON_BANNER].components
        self.assertEqual(b.section_title_symbol, "▍")
        self.assertEqual(b.section_title_gap, " ")
        self.assertEqual(b.title_numbering_template, None)
        self.assertEqual(b.role_separator, "  ·  ")
        self.assertEqual(b.date_mode, "tab_stop")
        self.assertEqual(b.date_prefix, "\t")
        self.assertEqual(b.bullet_symbol, "·")
        self.assertEqual(b.bullet_gap, " ")
        self.assertEqual(b.cert_label, "证书  ")
        self.assertEqual(b.cert_separator, "  ·  ")
        self.assertIsNone(b.side_label_transform)
        self.assertEqual(b.identity_text_align, "left")
        self.assertEqual(b.contact_align, "left")
        self.assertEqual(b.photo_align, "right")
        self.assertIs(b.photo_preserve_aspect, True)

        m = SKELETON_DEFAULT_STYLE[SKELETON_MINIMAL].components
        self.assertIsNone(m.section_title_symbol)
        self.assertEqual(m.title_numbering_template, "%02d")
        self.assertEqual(m.title_numbering_gap, "  ")
        self.assertEqual(m.role_separator, "  ")
        self.assertEqual(m.cert_label, "证书资质  ")

        s = SKELETON_DEFAULT_STYLE[SKELETON_SIDEBAR].components
        self.assertEqual(s.section_title_symbol, "▍")
        self.assertEqual(s.date_mode, "column")
        self.assertEqual(s.side_label_transform, "upper_ascii")
        self.assertEqual(s.identity_text_align, "center")
        self.assertEqual(s.photo_align, "center")
        self.assertEqual(s.contact_align, "left")
        self.assertIsNone(s.cert_label)

    def test_05b_identity_alignment_unknown_falls_back(self):
        from resume_generator.design import build_p1_spec
        spec = build_p1_spec()
        spec = replace(
            spec,
            header=replace(spec.header,
                           identity_alignment="unknown_future_mode"))
        st = resolve_spec_style(spec, SKELETON_BANNER)
        # 未知复合描述 → Profile（banner: left/right）+ 单次登记
        self.assertEqual(st.components.identity_text_align, "left")
        self.assertEqual(st.components.photo_align, "right")
        self.assertEqual(
            st.spec_fallbacks.count(
                "components.identity_alignment<not_defined>"), 1)

    def test_05c_p2_icon_title_capability_gap_registered(self):
        from resume_generator.design import build_p2_spec
        fb = resolve_spec_style(build_p2_spec(),
                                SKELETON_SIDEBAR).spec_fallbacks
        # P2 首选图标标题但图标资源是未定义缺口：只登记，渲染回退 ▍ 符号
        self.assertIn("components.section_icon<capability_gap>", fb)


# ===========================================================================
# 二、组件渲染（真实 DOCX）
# ===========================================================================

class TestComponentRendering(unittest.TestCase):

    # -- Test 6：SectionTitle -----------------------------------------------
    def test_06_section_title_symbol_and_numbering(self):
        # banner 旧路径：符号式
        d = build_document(fake_blocks(), SKELETON_BANNER, "tech_navy")
        t = find_para(d, "个人优势")
        self.assertTrue(t.text.startswith("▍ "))
        # sidebar 旧路径 _main_title 同样符号式
        d3 = build_document(fake_blocks(), SKELETON_SIDEBAR,
                            "manufacturing_steel")
        self.assertTrue(find_para(d3, "个人优势").text.startswith("▍ "))

        # minimal 旧路径：编号式 01/02…（编号 run 与文字 run 分离）
        d2 = build_document(fake_blocks(), SKELETON_MINIMAL, "minimal_ink")
        n1 = find_para(d2, "个人优势")
        self.assertTrue(n1.text.startswith("01  "))
        runs = n1.runs
        self.assertEqual(runs[0].text, "01  ")
        self.assertTrue(n1.text.endswith("个人优势"))

        # Spec 路径：自定义标题符号（尾空格是组装规则，不进 Spec）
        from resume_generator.design import build_p1_spec
        spec = build_p1_spec()
        spec = replace(
            spec,
            section=replace(spec.section,
                            prefix=Sourced.kv1("■", notes="测试自定义符号")))
        ds = build_document(fake_blocks(), design_spec=spec)
        self.assertTrue(find_para(ds, "个人优势").text.startswith("■ "))

    # -- Test 7：ExperienceHeader -------------------------------------------
    def test_07_experience_header_separator_and_date(self):
        # banner 旧路径："  ·  " 分隔 + \t 日期 + 右制表位
        d = build_document(fake_blocks(), SKELETON_BANNER, "tech_navy")
        hp = find_para(d, "某某有限公司")
        self.assertIn("某某有限公司  ·  测试实习生", hp.text)
        self.assertIn("\t2025.06 - 2025.09", hp.text)
        self.assertTrue(has_tab_stop(hp))

        # minimal 旧路径：双空格（无中点）
        d2 = build_document(fake_blocks(), SKELETON_MINIMAL, "minimal_ink")
        hp2 = find_para(d2, "某某有限公司")
        self.assertIn("某某有限公司  测试实习生", hp2.text)
        self.assertNotIn("·", hp2.text)

        # Spec P1：全角竖线（显式 Spec 值驱动）
        from resume_generator.design import build_p1_spec
        ds = build_document(fake_blocks(), design_spec=build_p1_spec())
        hps = find_para(ds, "某某有限公司")
        self.assertIn("｜测试实习生", hps.text)

        # sidebar 旧路径：column 模式同样以右缘制表位落地
        d3 = build_document(fake_blocks(), SKELETON_SIDEBAR,
                            "manufacturing_steel")
        hp3 = find_para(d3, "某某有限公司")
        self.assertIn("  ·  测试实习生", hp3.text)
        self.assertTrue(has_tab_stop(hp3))

    # -- Test 8：Bullet ------------------------------------------------------
    def test_08_bullet_prefix(self):
        d = build_document(fake_blocks(), SKELETON_BANNER, "tech_navy")
        self.assertTrue(
            find_para(d, "完成了一项虚构的测试工作").text.startswith("· "))
        # 侧栏技能 bullet 使用侧栏字号档，但前缀同一组件槽位
        d3 = build_document(fake_blocks(), SKELETON_SIDEBAR,
                            "manufacturing_steel")
        self.assertTrue(find_para(d3, "技能甲").text.startswith("· "))

        # Spec 路径：bullet.symbol="•"（Spec 明确值允许改变 Spec 路径输出）
        from resume_generator.design import build_p1_spec
        ds = build_document(fake_blocks(), design_spec=build_p1_spec())
        self.assertTrue(
            find_para(ds, "完成了一项虚构的测试工作").text.startswith("• "))

    # -- Test 9：Tag 现状锁定（纯文本 muted 行，无 chip 结构）---------------
    def test_09_tag_is_plain_muted_text_line(self):
        d = build_document(fake_blocks(), SKELETON_BANNER, "tech_navy")
        tp = find_para(d, "标签甲")
        # 当前能力：整串 tags 文本由数据层给出；Renderer 不组装分隔符
        self.assertIn("标签甲 · 标签乙", tp.text)
        pPr = tp._p.find(qn("w:pPr"))
        self.assertIsNotNone(pPr)
        # 无段落边框 / 无底纹（CSS 式 tag chip 属 implementation gap）
        self.assertIsNone(pPr.find(qn("w:pBdr")))
        self.assertIsNone(pPr.find(qn("w:shd")))
        for r in tp.runs:
            rPr = r._r.find(qn("w:rPr"))
            if rPr is not None:
                self.assertIsNone(rPr.find(qn("w:shd")))
                self.assertIsNone(rPr.find(qn("w:bdr")))

    # -- Test 10：Photo ------------------------------------------------------
    def test_10_photo_insertion_and_missing_behavior(self):
        with tempfile.TemporaryDirectory() as tmp:
            png = Path(tmp) / "id.png"
            png.write_bytes(_TINY_PNG)
            d = build_document(fake_blocks(str(png)), SKELETON_BANNER,
                               "tech_navy")
            # 等比插入：恰好 1 个 inline shape，落在右对齐照片段
            self.assertEqual(len(d.inline_shapes), 1)
            photo_paras = [p for p in all_paragraphs(d) if has_drawing(p)]
            self.assertEqual(len(photo_paras), 1)
            self.assertEqual(jc_of_para(photo_paras[0]), "right")

            d3 = build_document(fake_blocks(str(png)), SKELETON_SIDEBAR,
                                "manufacturing_steel")
            self.assertEqual(len(d3.inline_shapes), 1)
            photo_paras3 = [p for p in all_paragraphs(d3) if has_drawing(p)]
            self.assertEqual(jc_of_para(photo_paras3[0]), "center")

        # 缺图：不插入、不抛异常、不产出 drawing
        d0 = build_document(fake_blocks(None), SKELETON_BANNER, "tech_navy")
        self.assertEqual(len(d0.inline_shapes), 0)
        self.assertFalse(
            any(has_drawing(p) for p in all_paragraphs(d0)))

        # preserve_aspect 槽位：三骨架默认 True，Spec 路径同样 True
        from resume_generator.design import build_p1_spec
        st = resolve_spec_style(build_p1_spec(), SKELETON_BANNER)
        self.assertIs(st.components.photo_preserve_aspect, True)

    # -- Test 11：AccentBar / Rule ------------------------------------------
    def test_11_accent_bar_bottom_border(self):
        from resume_generator.render_style import build_palette_for_skeleton
        rp = build_palette_for_skeleton(
            SKELETON_BANNER, resolve_palette("tech_navy", SKELETON_BANNER))
        st = SKELETON_DEFAULT_STYLE[SKELETON_BANNER]
        d = build_document(fake_blocks(), SKELETON_BANNER, "tech_navy")
        bars = [bottom_border(p) for p in d.paragraphs]
        bars = [b for b in bars if b is not None]
        # 色条：thickness 来自 Spacing（bar_sz=18），颜色来自 Palette.accent
        self.assertIn({"sz": str(st.spacing.bar_sz),
                       "color": rp.accent.lstrip("#"),
                       "val": "single"}, bars)

        # minimal rule：rule_sz Profile
        st2 = SKELETON_DEFAULT_STYLE[SKELETON_MINIMAL]
        d2 = build_document(fake_blocks(), SKELETON_MINIMAL, "minimal_ink")
        bars2 = [b for b in (bottom_border(p) for p in d2.paragraphs)
                 if b is not None]
        self.assertIn(str(st2.spacing.rule_sz), [b["sz"] for b in bars2])

    # -- Test 12：ContactBlock ----------------------------------------------
    def test_12_contact_block_plain_paragraphs(self):
        d = build_document(fake_blocks(), SKELETON_BANNER, "tech_navy")
        cp = find_para(d, "zhangsan@example.com")
        # 当前实现：普通文本段落（图标/label-value 结构属 implementation gap）
        self.assertEqual(cp.text, "zhangsan@example.com")
        self.assertFalse(has_drawing(cp))
        pPr = cp._p.find(qn("w:pPr"))
        self.assertIsNone(pPr.find(qn("w:pBdr")))
        # contact_align="left" 不写 w:jc（OOXML 默认），不产生多余 XML
        self.assertIsNone(jc_of_para(cp))

        # sidebar：联系方式段同样无 w:jc；照片/姓名/意向居中
        d3 = build_document(fake_blocks(), SKELETON_SIDEBAR,
                            "manufacturing_steel")
        cp3 = find_para(d3, "zhangsan@example.com")
        self.assertIsNone(jc_of_para(cp3))
        self.assertEqual(jc_of_para(find_para(d3, "张三")), "center")
        self.assertEqual(jc_of_para(find_para(d3, "测试岗位")), "center")

    # -- Test 13：证书行（banner/minimal 标签与 join 走组件槽位）------------
    def test_13_cert_line_labels(self):
        d = build_document(fake_blocks(), SKELETON_BANNER, "tech_navy")
        p = find_para(d, "虚构证书甲")
        self.assertIn("证书  虚构证书甲  ·  虚构证书乙", p.text)

        d2 = build_document(fake_blocks(), SKELETON_MINIMAL, "minimal_ink")
        p2 = find_para(d2, "虚构证书甲")
        self.assertIn("证书资质  虚构证书甲  ·  虚构证书乙", p2.text)

    # -- Test 14：侧栏分组标签大小写变换 ------------------------------------
    def test_14_side_label_ascii_uppercase(self):
        from docx import Document
        from resume_generator.render_style import build_palette_for_skeleton
        from resume_generator.skeletons import _side_label

        st = replace(
            SKELETON_DEFAULT_STYLE[SKELETON_SIDEBAR],
            palette=build_palette_for_skeleton(
                SKELETON_SIDEBAR,
                resolve_palette(None, SKELETON_SIDEBAR)))
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        _side_label(cell, "skills", st)
        _side_label(cell, "技能", st)
        texts = [p.text for p in cell.paragraphs]
        # upper_ascii 机制：ASCII 标签全大写；中文原样
        self.assertIn("SKILLS", texts)
        self.assertIn("技能", texts)

        # banner/minimal 不消费该槽位（None）→ 若启用也不会改写
        self.assertIsNone(
            SKELETON_DEFAULT_STYLE[SKELETON_BANNER]
            .components.side_label_transform)


# ===========================================================================
# 二之二、P3 symbol_hairline 标题契约（Phase 2B-1 Step 4 / P-1）
# ===========================================================================

class TestP3SymbolTitleRendering(unittest.TestCase):
    """P3 minimal 板块标题：▌ 符号 run（accent #0A6B3C）+ 标题文字 run（ink）。

    Spec：section.title_style="symbol_hairline"、prefix="▌"、
    symbol_color_role="accent"、title_color_role="ink"。
    DOCX 全部在内存构建（build_document 不 save），无项目目录写入。
    """

    # fake_blocks 实际产出的 minimal 板块（projects=[] 故无「项目经历」）
    TITLES = ("个人优势", "教育经历", "实习经历", "校园经历",
              "专业技能与证书")

    def _title_paragraphs(self):
        from resume_generator.design import build_p3_spec
        doc = build_document(fake_blocks(), design_spec=build_p3_spec())
        wanted = {"▌ " + t for t in self.TITLES}
        return [p for p in all_paragraphs(doc) if p.text in wanted]

    def test_p3_title_uses_symbol_prefix(self):
        paras = self._title_paragraphs()
        # 五个板块标题全部以 ▌ 起头，且各只出现一次 ▌
        self.assertEqual(len(paras), len(self.TITLES))
        for p in paras:
            self.assertTrue(p.text.startswith("▌ "), msg=p.text)
            self.assertEqual(p.text.count("▌"), 1, msg=p.text)
            self.assertEqual(
                len([r for r in p.runs if "▌" in r.text]), 1, msg=p.text)

    def test_p3_title_has_no_numbering(self):
        # 标题段不允许任何数字（旧 minimal 路径的 01/02/03 编号不得出现）
        for p in self._title_paragraphs():
            for token in ("01", "02", "03"):
                self.assertNotIn(token, p.text, msg=p.text)
            self.assertNotRegex(p.text, r"\d", msg=p.text)
            self.assertEqual(len(p.runs), 2, msg=p.text)

    def test_p3_symbol_colored_accent(self):
        # 含 ▌ 的唯一 run：w:color=0A6B3C（symbol_color_role=accent）
        for p in self._title_paragraphs():
            symbol_runs = [r for r in p.runs if "▌" in r.text]
            self.assertEqual(len(symbol_runs), 1, msg=p.text)
            self.assertEqual(run_color(symbol_runs[0]), "0A6B3C",
                             msg=p.text)

    def test_p3_title_text_remains_ink(self):
        # 标题文字 run（唯一不含 ▌ 的 run）：w:color=000000（title_color_role=ink）
        for p in self._title_paragraphs():
            text_runs = [r for r in p.runs if "▌" not in r.text]
            self.assertEqual(len(text_runs), 1, msg=p.text)
            self.assertEqual(run_color(text_runs[0]), "000000", msg=p.text)


# ===========================================================================
# 二之三、P3 Header 单列结构契约（Phase 2B-2 Step 5 / P-2）
# ===========================================================================

class TestP3HeaderSingleColumn(unittest.TestCase):
    """P-2：P3 spec.photo.enabled=False / photo_zone="none" → Header 单列。

    旧路径（无 spec）与 P1/P2 spec 路径均不经此分支，行为字节级不变
    （smoke_layout_kit.py 仍 37983/37933/37974 byte-perfect）。
    DOCX 全部在内存构建，无项目目录写入。
    """

    def _p3_doc(self):
        from resume_generator.design import build_p3_spec
        return build_document(fake_blocks(), design_spec=build_p3_spec())

    @staticmethod
    def _grid_col_widths_twips(table):
        """返回表 tblGrid 中各 w:gridCol 的 w:w（twips/dxa）列表。"""
        grid = table._tbl.find(qn("w:tblGrid"))
        cols = grid.findall(qn("w:gridCol"))
        return [int(c.get(qn("w:w"))) for c in cols]

    # -- 1. 单列 tblGrid ---------------------------------------------------
    def test_p3_header_has_one_column(self):
        doc = self._p3_doc()
        widths = self._grid_col_widths_twips(doc.tables[0])
        self.assertEqual(len(widths), 1,
                         f"P3 Header 应单列，实际 {len(widths)} 列")

    # -- 2. 无右侧空照片格 / 无 insert_photo 痕迹 --------------------------
    def test_p3_header_no_photo_cell(self):
        doc = self._p3_doc()
        header_row = doc.tables[0].rows[0]
        # 单格 Header：无第二个单元格
        self.assertEqual(len(header_row.cells), 1)
        # P3 spec 路径 blocks.photo_path=None 且不再调用 insert_photo
        self.assertEqual(len(doc.inline_shapes), 0)
        self.assertFalse(any(has_drawing(p) for p in all_paragraphs(doc)))

    # -- 3. 单列宽度 = usable_width（≈ 18.319 cm）-------------------------
    def test_p3_header_text_uses_full_width(self):
        from docx.shared import Cm
        from resume_generator.layout_kit import usable_width_cm

        doc = self._p3_doc()
        widths = self._grid_col_widths_twips(doc.tables[0])
        # P3 margins: left=1.499 right=1.182（presets.py:566-567）→ usable=18.319cm
        expected = int(Cm(usable_width_cm(1.499, 1.182)).twips)
        self.assertEqual(widths[0], expected,
                         f"单列宽度应={expected} twips，实际={widths[0]}")
        # 防御：换算回 cm ≈ 18.319（1 cm = 566.929 twips）
        cm_actual = widths[0] / 566.929
        self.assertAlmostEqual(cm_actual, 18.319, delta=0.005)

    # -- 4. 文本内容 / 字号 / 字重 / 颜色 与原 P3 一致 ----------------------
    def test_p3_header_content_unchanged(self):
        import math

        def half_pt(v):
            """python-docx 把 pt 截断到 w:sz 半磅整数网格后的回读值。"""
            return math.floor(v * 2) / 2

        def is_bold(run):
            """读 w:b 的 w:val：absent/1/true = 加粗；0/false = 不加粗。"""
            rPr = run._r.find(qn("w:rPr"))
            if rPr is None:
                return False
            b = rPr.find(qn("w:b"))
            if b is None:
                return False
            val = b.get(qn("w:val"))
            return val is None or val in ("1", "true")

        doc = self._p3_doc()

        # name：17.04pt → 半磅 17.0；bold (w:b)；ink = 000000
        p_name = find_para(doc, "张三")
        name_run = next(r for r in p_name.runs if r.text == "张三")
        self.assertAlmostEqual(name_run.font.size.pt,
                               half_pt(17.04), delta=0.005)
        self.assertTrue(is_bold(name_run),
                        "name run 应加粗（w:b val=1/absent）")
        self.assertEqual(run_color(name_run), "000000")

        # intent：11pt；bold（build_minimal 硬编码 bold=True）；accent = 0A6B3C
        p_intent = find_para(doc, "测试岗位")
        intent_run = next(r for r in p_intent.runs if r.text == "测试岗位")
        self.assertAlmostEqual(intent_run.font.size.pt,
                               half_pt(11.0), delta=0.005)
        self.assertTrue(is_bold(intent_run),
                        "intent run 应加粗（build_minimal 硬编码）")
        self.assertEqual(run_color(intent_run), "0A6B3C")

        # contact：8.52pt → 半磅 8.5；不 bold（w:b val=0）；muted = 444444
        p_contact = find_para(doc, "138-0000-0000")
        contact_run = next(r for r in p_contact.runs
                           if r.text == "138-0000-0000")
        self.assertAlmostEqual(contact_run.font.size.pt,
                               half_pt(8.52), delta=0.005)
        self.assertFalse(is_bold(contact_run),
                         "contact run 不应加粗（w:b val=0）")
        self.assertEqual(run_color(contact_run), "444444")

    # -- 5. P1/P2 spec 路径 Header 列数不变（保护边界）---------------------
    def test_p1_p2_header_unchanged(self):
        from resume_generator.design import build_p1_spec, build_p2_spec

        d1 = build_document(fake_blocks(), design_spec=build_p1_spec())
        self.assertEqual(len(self._grid_col_widths_twips(d1.tables[0])),
                         2, "P1 banner Header 应保持 2 列")
        d2 = build_document(fake_blocks(), design_spec=build_p2_spec())
        self.assertEqual(len(self._grid_col_widths_twips(d2.tables[0])),
                         2, "P2 sidebar 表应保持 2 列")

    # -- 6. 旧路径（无 spec）minimal Header 仍 2 列含空照片格 ------------
    def test_p3_legacy_path_keeps_two_columns(self):
        d = build_document(fake_blocks(), SKELETON_MINIMAL, "minimal_ink")
        widths = self._grid_col_widths_twips(d.tables[0])
        self.assertEqual(len(widths), 2, "旧路径 minimal Header 应保持 2 列")
        # 右列空照片格仍存在（无 drawing 但单元格在）
        self.assertEqual(len(d.tables[0].rows[0].cells), 2)
        self.assertEqual(len(d.inline_shapes), 0)


# ===========================================================================
# 二之四、P3 accent_line 关闭契约（Phase 2B-1 Step 6 / P-3）
# ===========================================================================

class TestP3AccentLineDisabled(unittest.TestCase):
    """P-3：P3 spec.header.accent_line.enabled=False → minimal Header 顶部
    不绘制强调线（不创建空段 / 不写 w:pBdr / 不写 rule_sz）。

    旧路径（无 spec）与 P1/P2 spec 路径均不经 build_minimal，行为字节级不变。
    DOCX 全部在内存构建，无项目目录写入。
    """

    def _p3_doc(self):
        from resume_generator.design import build_p3_spec
        return build_document(fake_blocks(), design_spec=build_p3_spec())

    # -- 1. P3 不包含 Header 顶部强调线段（无 accent 色 w:pBdr/bottom）-------
    def test_p3_no_header_accent_rule(self):
        doc = self._p3_doc()
        # P3 accent 色 = #0A6B3C；rule 段若存在必含此色
        for p in doc.paragraphs:
            bb = bottom_border(p)
            if bb is None:
                continue
            self.assertNotEqual(bb.get("color"), "0A6B3C",
                                f"P3 Header 顶部强调线不应存在：{p.text!r}")

    # -- 2. P3 Header 后直接接板块标题（无空 rule 段）-----------------------
    def test_p3_no_empty_rule_paragraph(self):
        doc = self._p3_doc()
        # rule 段是空文本段（p.text == ""）紧跟 Header 表格之后
        # P3 应直接进入「个人优势」标题（fake_blocks 含 summary）
        body_paras = [p for p in doc.paragraphs if p.text != ""]
        # 第一个 body 段应是「▌ 个人优势」标题，而非空 rule 段
        first = body_paras[0]
        self.assertTrue(first.text.startswith("▌ 个人优势"),
                        f"P3 应直接进入板块标题，实际首段={first.text!r}")

    # -- 3. P1/P2 spec 路径强调线行为不变 --------------------------------
    def test_p1_p2_accent_line_unchanged(self):
        from resume_generator.design import build_p1_spec, build_p2_spec
        # P1 banner：accent_line.enabled=True → bar 段存在且含 accent 色
        d1 = build_document(fake_blocks(), design_spec=build_p1_spec())
        p1_bars = [bottom_border(p) for p in d1.paragraphs
                   if bottom_border(p) is not None]
        self.assertTrue(any(b.get("color") == "E31937" for b in p1_bars),
                        "P1 banner accent_line 仍应绘制")
        # P2 sidebar：无 rule/bar 段（build_sidebar 不创建顶部线）
        d2 = build_document(fake_blocks(), design_spec=build_p2_spec())
        # sidebar 顶部线在 side_label 段，非 Header 顶部 accent_line
        # P2 accent_line.enabled=False，不应产生独立 bar 段
        # （side_label 的 bottom border 是另一个语义，不在本断言范围）
        self.assertEqual(d2.tables[0].rows[0].cells[0].paragraphs[0].text,
                         "", "P2 sidebar 不受 P-3 影响")

    # -- 4. 旧路径 minimal rule 仍存在 ------------------------------------
    def test_legacy_minimal_rule_unchanged(self):
        d = build_document(fake_blocks(), SKELETON_MINIMAL, "minimal_ink")
        bars = [bottom_border(p) for p in d.paragraphs
                if bottom_border(p) is not None]
        self.assertTrue(len(bars) >= 1, "旧路径 minimal 应保留 rule 段")


# ===========================================================================
# 二之五、P3 latin_font 落地契约（Phase 2B-1 Step 7 / P-4）
# ===========================================================================

class TestP3LatinFontRendering(unittest.TestCase):
    """P-4：P3 typography.latin_font="ArialMT" 真正进入 DOCX。

    期望 rFonts：w:eastAsia="微软雅黑"（font_family Microsoft YaHei 经
    字体令牌映射）、w:ascii="ArialMT"、w:hAnsi="ArialMT"。
    P1/P2 与旧路径 latin_font=None → 三槽位同写 微软雅黑（字节保值）。
    DOCX 全部在内存构建，无项目目录写入。
    """

    @staticmethod
    def _run_rfonts(run):
        rPr = run._r.find(qn("w:rPr"))
        if rPr is None:
            return None
        return rPr.find(qn("w:rFonts"))

    def _rfonts_of(self, run):
        rf = self._run_rfonts(run)
        self.assertIsNotNone(rf, "run 缺少 w:rFonts")
        return {
            "eastAsia": rf.get(qn("w:eastAsia")),
            "ascii": rf.get(qn("w:ascii")),
            "hAnsi": rf.get(qn("w:hAnsi")),
        }

    def test_p3_run_rfonts_split_cjk_and_latin(self):
        from resume_generator.design import build_p3_spec
        doc = build_document(fake_blocks(), design_spec=build_p3_spec())
        # 中文 run（姓名）与拉丁 run（邮箱）都显式带字体，槽位分离
        name_run = next(r for r in find_para(doc, "张三").runs
                        if r.text == "张三")
        self.assertEqual(self._rfonts_of(name_run),
                         {"eastAsia": "微软雅黑",
                          "ascii": "ArialMT", "hAnsi": "ArialMT"})

        mail_run = next(r for r in find_para(doc, "zhangsan@example.com").runs
                        if "zhangsan" in r.text)
        self.assertEqual(self._rfonts_of(mail_run),
                         {"eastAsia": "微软雅黑",
                          "ascii": "ArialMT", "hAnsi": "ArialMT"})

        # 标题符号 run / 正文 bullet run 同样分离
        sym_title = next(p for p in all_paragraphs(doc)
                         if p.text.startswith("▌ "))
        for r in sym_title.runs:
            self.assertEqual(self._rfonts_of(r)["ascii"], "ArialMT")
            self.assertEqual(self._rfonts_of(r)["hAnsi"], "ArialMT")
            self.assertEqual(self._rfonts_of(r)["eastAsia"], "微软雅黑")

    def test_p3_normal_style_rfonts_split(self):
        from resume_generator.design import build_p3_spec
        doc = build_document(fake_blocks(), design_spec=build_p3_spec())
        rFonts = doc.styles["Normal"].element.rPr.rFonts
        self.assertEqual(rFonts.get(qn("w:eastAsia")), "微软雅黑")
        self.assertEqual(rFonts.get(qn("w:ascii")), "ArialMT")
        self.assertEqual(rFonts.get(qn("w:hAnsi")), "ArialMT")

    def test_p3_resolved_style_carries_latin_font(self):
        from resume_generator.design import build_p3_spec
        st = resolve_spec_style(build_p3_spec(), SKELETON_MINIMAL)
        self.assertEqual(st.latin_font, "ArialMT")
        self.assertEqual(st.font_family, "微软雅黑")

    def test_p1_p2_latin_font_none_keeps_single_family(self):
        from resume_generator.design import build_p1_spec, build_p2_spec
        # P1/P2 spec 未声明 latin_font → None，三槽位同写
        self.assertIsNone(
            resolve_spec_style(build_p1_spec(), SKELETON_BANNER).latin_font)
        self.assertIsNone(
            resolve_spec_style(build_p2_spec(), SKELETON_SIDEBAR).latin_font)

        d1 = build_document(fake_blocks(), design_spec=build_p1_spec())
        # banner 姓名在表格单元格内
        name_run = None
        for p in all_paragraphs(d1):
            for r in p.runs:
                if r.text == "张三":
                    name_run = r
                    break
            if name_run is not None:
                break
        self.assertIsNotNone(name_run)
        self.assertEqual(self._rfonts_of(name_run),
                         {"eastAsia": "微软雅黑",
                          "ascii": "微软雅黑", "hAnsi": "微软雅黑"})

    def test_legacy_path_rfonts_unchanged(self):
        # 旧路径三骨架：latin_font=None → 三槽位同写 微软雅黑
        for sk, palette in ((SKELETON_BANNER, "tech_navy"),
                            (SKELETON_MINIMAL, "minimal_ink"),
                            (SKELETON_SIDEBAR, "manufacturing_steel")):
            doc = build_document(fake_blocks(), sk, palette)
            run = None
            for p in all_paragraphs(doc):
                for r in p.runs:
                    if r.text == "张三":
                        run = r
                        break
                if run is not None:
                    break
            self.assertIsNotNone(run, f"{sk} 缺少姓名 run")
            self.assertEqual(self._rfonts_of(run),
                             {"eastAsia": "微软雅黑",
                              "ascii": "微软雅黑", "hAnsi": "微软雅黑"},
                             msg=sk)


# ===========================================================================
# 二之六、P3 divider 契约（Phase 2B-1 Step 8 / P-5；2B-4 Step 3 对齐 Golden Sample）
# ===========================================================================

class TestP3DividerGoldenAligned(unittest.TestCase):
    """P-5：P3 section.divider.weight=2.25pt（Golden Sample CL-02）→
    每个板块标题下生成独立 hairline 段（pBdr 底边框，sz=18、绿 #007A37）。

    标题段自身不附 pBdr（divider 是独立空段）；标题段
    section_before/title_after 保留。
    P1 spec 仍为 0pt（Spec 明文「0pt 表达禁用」）→ banner 板块 hairline
    一并消失，但 accent 色条（header.accent_line，另一语义）保留。
    P2 spec=0.5pt 与旧路径=6 均 >0，divider 行为字节级不变。
    """

    def _p3_doc(self):
        from resume_generator.design import build_p3_spec
        return build_document(fake_blocks(), design_spec=build_p3_spec())

    @staticmethod
    def _all_bottom_borders(doc):
        out = []
        for p in all_paragraphs(doc):
            bb = bottom_border(p)
            if bb is not None:
                out.append((p, bb))
        return out

    # -- 1. P3 divider 底边框：sz=18（2.25pt）、绿 007A37，且全部 pBdr 均为该语义 --
    def test_p3_divider_borders_golden_values(self):
        doc = self._p3_doc()
        borders = [{**b, "color": b["color"].upper()}
                   for _, b in self._all_bottom_borders(doc)]
        self.assertIn({"sz": "18", "color": "007A37", "val": "single"},
                      borders)
        # P3 无 rule（P-3 已移除）：全部底边框均应属 divider 语义
        self.assertTrue(
            all(b["sz"] == "18" and b["color"] == "007A37" for b in borders),
            msg=f"存在非 divider 底边框：{borders}")

    # -- 2. P3 标题段本身不附 pBdr ---------------------------------------
    def test_p3_title_paragraph_has_no_pBdr(self):
        doc = self._p3_doc()
        for p in all_paragraphs(doc):
            if p.text.startswith("▌ "):
                pPr = p._p.find(qn("w:pPr"))
                self.assertIsNotNone(pPr)
                self.assertIsNone(pPr.find(qn("w:pBdr")), msg=p.text)

    # -- 3. 每个板块标题后紧邻 divider 空段（add_hairline：pBdr sz=18 绿）----
    def test_p3_divider_paragraph_after_every_title(self):
        doc = self._p3_doc()
        paras = doc.paragraphs
        title_count = 0
        for i, p in enumerate(paras):
            if not p.text.startswith("▌ "):
                continue
            title_count += 1
            nxt = paras[i + 1]
            self.assertEqual(nxt.text, "",
                             msg=f"标题 {p.text!r} 后不是 divider 空段")
            bb = bottom_border(nxt)
            self.assertIsNotNone(bb, msg=p.text)
            self.assertEqual(bb["sz"], "18", msg=p.text)
            self.assertEqual(bb["color"].upper(), "007A37", msg=p.text)
        self.assertGreater(title_count, 0)

    # -- 4. 标题段自身段前/段后间距保留（未被误删）-------------------------
    def test_p3_title_spacing_preserved(self):
        from docx.oxml.ns import qn as _qn
        doc = self._p3_doc()
        title = next(p for p in doc.paragraphs if p.text.startswith("▌ "))
        pPr = title._p.find(_qn("w:pPr"))
        spacing = pPr.find(_qn("w:spacing"))
        self.assertIsNotNone(spacing)
        # section_before 由 Spec 10pt 解析 → w:before=200（1/20 pt）
        self.assertEqual(spacing.get(_qn("w:before")), "200")
        # title_after 沿用 minimal Profile 1pt → 20
        self.assertEqual(spacing.get(_qn("w:after")), "20")

    # -- 5. RenderSpacing：P3 hairline_sz=18（2.25pt×8）、hairline 色=007A37 --
    def test_p3_resolved_hairline_sz_18(self):
        from resume_generator.design import build_p3_spec
        st = resolve_spec_style(build_p3_spec(), SKELETON_MINIMAL)
        self.assertEqual(st.spacing.hairline_sz, 18)
        self.assertEqual(st.palette.hairline, "#007A37")

    # -- 6. P1 spec：divider 0pt 不绘制，但 accent 色条保留 ----------------
    def test_p1_spec_divider_off_but_accent_bar_kept(self):
        from resume_generator.design import build_p1_spec
        d1 = build_document(fake_blocks(), design_spec=build_p1_spec())
        borders = [b for _, b in self._all_bottom_borders(d1)]
        colors = {b["color"].upper() for b in borders}
        # 板块 hairline（#DFDFDF）随 0pt 禁用消失
        self.assertNotIn("DFDFDF", colors)
        # header accent_line 色条保留，语义独立；厚度由 Spec 2.16pt×8=17
        self.assertIn("E31937", colors)
        self.assertIn({"sz": "17", "color": "E31937", "val": "single"},
                      [{**b, "color": b["color"].upper()} for b in borders])

    # -- 7. P2 spec：0.5pt divider 仍附在 sidebar 标题段 ------------------
    def test_p2_spec_divider_remains(self):
        from resume_generator.design import build_p2_spec
        d2 = build_document(fake_blocks(), design_spec=build_p2_spec())
        borders = [{**b, "color": b["color"].upper()}
                   for _, b in self._all_bottom_borders(d2)]
        # 0.5pt → sz = round(0.5*8) = 4；hairline 色 DDE2E8（Spec 原值大写）
        self.assertIn({"sz": "4", "color": "DDE2E8", "val": "single"},
                      borders)

    # -- 8. 旧路径三骨架 divider 全部保留（smoke 之外的结构守卫）----------
    def test_legacy_paths_dividers_unchanged(self):
        d_b = build_document(fake_blocks(), SKELETON_BANNER, "tech_navy")
        b_borders = [b for _, b in self._all_bottom_borders(d_b)]
        # banner：hairline sz=6 若干 + accent bar sz=18
        self.assertIn("6", {b["sz"] for b in b_borders})
        self.assertIn("18", {b["sz"] for b in b_borders})

        d_m = build_document(fake_blocks(), SKELETON_MINIMAL, "minimal_ink")
        m_sizes = {b["sz"] for _, b in self._all_bottom_borders(d_m)}
        # minimal：rule sz=12 + 板块 hairline sz=6
        self.assertIn("12", m_sizes)
        self.assertIn("6", m_sizes)

        d_s = build_document(fake_blocks(), SKELETON_SIDEBAR,
                             "manufacturing_steel")
        s_sizes = {b["sz"] for _, b in self._all_bottom_borders(d_s)}
        # sidebar：侧栏标签 sz=8 + 主栏标题 hairline sz=6
        self.assertIn("8", s_sizes)
        self.assertIn("6", s_sizes)


# ===========================================================================
# 三、源码守卫 + 跨套件回归
# ===========================================================================

class TestComponentSourceGuard(unittest.TestCase):

    def test_15_no_component_literals_in_skeletons(self):
        """组件结构字面量必须只存在于 render_style 的 Profile / Resolver。

        tokenize 剥除注释后扫描 skeletons.py：
          - 禁止 STRING 字面量："· "、"  ·  "、"▍ "、"证书  "、"证书资质  "、
            "\\t"，以及任何含 ":02d" 的格式串；
          - 禁止再出现 role_separator 形参/变量名；
          - 段落对齐枚举只允许出现在 _apply_align 内部；
          - .upper() 只允许出现在 _side_label 的 transform 分支内。
        """
        src_path = SRC / "resume_generator" / "skeletons.py"
        raw = src_path.read_text(encoding="utf-8")
        lines = raw.splitlines()

        def span(func_name):
            start = next(i for i, ln in enumerate(lines)
                         if ln.startswith(f"def {func_name}("))
            end = next((i for i in range(start + 1, len(lines))
                        if lines[i] and not lines[i][0].isspace()
                        and not lines[i].startswith("#")), len(lines))
            return set(range(start, end))

        align_span = span("_apply_align")
        label_span = span("_side_label")

        # 1) STRING 字面量扫描
        forbidden = {"· ", "  ·  ", "▍ ", "证书  ", "证书资质  ", "\t"}
        with tokenize.open(src_path) as f:
            toks = list(tokenize.generate_tokens(
                io.StringIO(f.read()).readline))
        for tok in toks:
            if tok.type != tokenize.STRING:
                continue
            try:
                val = eval(tok.string, {"__builtins__": {}})  # noqa: S307
            except Exception:  # noqa: BLE001 - f-string 片段等跳过
                continue
            if not isinstance(val, str):
                continue
            self.assertNotIn(
                val, forbidden,
                msg=f"skeletons.py:{tok.start[0]} 组件字符串回迁：{val!r}")
            self.assertNotIn(
                ":02d", val,
                msg=f"skeletons.py:{tok.start[0]} 编号格式应在 RenderComponents")

        # 2) role_separator 形参/实参彻底退出骨架（允许 cp.role_separator 读取）
        for i, ln in enumerate(lines, 1):
            compact = ln.replace(" ", "")
            self.assertFalse(
                "role_separator:" in compact or "role_separator=" in compact,
                msg=f"skeletons.py:{i} role_separator 不得再作形参/实参，"
                    f"只能读 components.role_separator")

        # 3) 对齐枚举只在 _apply_align
        for i, ln in enumerate(lines):
            if "WD_ALIGN_PARAGRAPH.RIGHT" in ln \
                    or "WD_ALIGN_PARAGRAPH.CENTER" in ln:
                self.assertIn(i, align_span,
                              msg=f"skeletons.py:{i + 1} 对齐枚举必须走 "
                                  f"_apply_align(style.components.*)")

        # 4) .upper() 只在 _side_label 的 transform 分支
        for i, ln in enumerate(lines):
            if ".upper()" in ln:
                self.assertIn(i, label_span,
                              msg=f"skeletons.py:{i + 1} .upper() 必须由 "
                                  f"side_label_transform 控制")

    def test_16_render_style_owns_component_profiles(self):
        """正向守卫：三个 Profile 必须在 render_style.py 内集中登记。"""
        rs = (SRC / "resume_generator" / "render_style.py").read_text(
            encoding="utf-8")
        self.assertIn("class RenderComponents", rs)
        self.assertIn("_DEFAULT_COMPONENTS", rs)
        self.assertIn("_resolve_components_for_spec", rs)
        self.assertIn("components: Optional[RenderComponents]", rs)

    def test_17_typography_regression_suite(self):
        proc = subprocess.run(
            [sys.executable, str(ROOT / "tests" /
                                 "test_renderer_typography.py")],
            cwd=str(ROOT), capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0,
                         msg=f"typography 回归失败：\n{proc.stdout}\n"
                             f"{proc.stderr}")

    def test_18_spacing_regression_suite(self):
        proc = subprocess.run(
            [sys.executable, str(ROOT / "tests" /
                                 "test_renderer_spacing.py")],
            cwd=str(ROOT), capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0,
                         msg=f"spacing 回归失败：\n{proc.stdout}\n"
                             f"{proc.stderr}")

    def test_19_palette_regression_suite(self):
        proc = subprocess.run(
            [sys.executable, str(ROOT / "tests" /
                                 "test_renderer_palette.py")],
            cwd=str(ROOT), capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0,
                         msg=f"palette 回归失败：\n{proc.stdout}\n"
                             f"{proc.stderr}")

    def test_20_stability_regression_suite(self):
        proc = subprocess.run(
            [sys.executable, str(ROOT / "tests" /
                                 "test_renderer_stability.py")],
            cwd=str(ROOT), capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0,
                         msg=f"stability 回归失败：\n{proc.stdout}\n"
                             f"{proc.stderr}")


if __name__ == "__main__":
    unittest.main(verbosity=2)
