# -*- coding: utf-8 -*-
"""test_renderer_typography.py — Phase 3B-1 Typography 接线回归。

验证两件事：
    A. 旧路径（build_document(blocks, skeleton_id)）排版参数与改造前
       审计原值逐 pt 一致（视觉保值）。
    B. Spec 路径（build_document(blocks, design_spec=...)）的
       字体/字号/行距/字重确实来自 TypographySpec；undetermined 字段
       回退骨架 Profile 且在 spec_fallbacks 登记（不编造）。

数据全部虚构（张三）。用法：python tests/test_renderer_typography.py
"""
from __future__ import annotations

import re
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from docx.oxml.ns import qn  # noqa: E402

from resume_generator.layout_kit import BulletBlock, ResumeBlocks  # noqa: E402
from resume_generator.render_style import (  # noqa: E402
    SKELETON_DEFAULT_STYLE,
    resolve_font_family,
    resolve_style_for_spec,
)
from resume_generator.skeletons import build_document, resolve_palette  # noqa: E402


def fake_blocks():
    return ResumeBlocks(
        name="张三",
        intent="测试岗位",
        contact_lines=["138-0000-0000", "zhangsan@example.com"],
        summary="虚构的个人优势简介，用于排版接线测试。",
        education_lines=["XX大学  测试专业（本科）"],
        internships=[
            BulletBlock(
                title="某某有限公司",
                role="测试实习生",
                meta="2025.06 - 2025.09",
                tags="标签甲 · 标签乙",
                bullets=["完成了一项虚构的测试工作。"],
            )
        ],
        projects=[],
        campus=["虚构校园经历甲"],
        skills=["技能甲", "技能乙"],
        certs=["虚构证书甲"],
    )


def all_paragraphs(doc):
    ps = list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                ps.extend(cell.paragraphs)
    return ps


def find_para(doc, needle):
    for p in all_paragraphs(doc):
        if needle in p.text:
            return p
    raise AssertionError(f"找不到包含 {needle!r} 的段落")


def run_sizes(p):
    """段落内有文字的 run → (文本, pt 字号)。"""
    out = []
    for r in p.runs:
        if r.text and r.font.size is not None:
            out.append((r.text, round(r.font.size.pt, 4)))
    return out


def first_run_font(p):
    r = p.runs[0]
    rfonts = r._r.find(qn("w:rPr")).find(qn("w:rFonts"))
    return rfonts.get(qn("w:ascii")), rfonts.get(qn("w:eastAsia"))


def size_of_para(doc, needle, run_index=0):
    p = find_para(doc, needle)
    sized = [(t, s) for t, s in run_sizes(p)]
    return sized[run_index][1], p


# 旧路径审计原值（3B-1 改造前硬编码，Profile 必须逐值保值）
LEGACY = {
    "banner_card": {
        "name": 18.0, "intent": 10.5, "contact": 9.0, "title": 12.0,
        "org": 10.5, "role": 10.0, "body": 9.5, "meta": 9.0, "tags": 8.5,
        "line_body": 1.08, "line_title": 1.0, "line_exp": 1.05,
    },
    "single_column_minimal": {
        "name": 20.0, "intent": 11.0, "contact": 9.0, "title": 12.0,
        "org": 10.5, "role": 10.0, "body": 9.5, "meta": 9.0, "tags": 8.5,
        "line_body": 1.06, "line_title": 1.0, "line_exp": 1.05,
    },
    "two_column_sidebar": {
        "name": 14.0, "intent": 8.5, "contact": 8.5, "title": 12.0,
        "org": 10.5, "role": 10.0, "body": 9.5, "meta": 9.0,
        "side_body": 8.5, "line_body": 1.08, "line_bullet": 1.06,
        "line_side": 1.05, "line_title": 1.0,
    },
}


class TestTypographyLegacyPreservation(unittest.TestCase):
    """A. 旧路径视觉保值：Profile 即审计原值，DOCX 取值一致。"""

    def test_01_profiles_equal_audited_values(self):
        for sk, want in LEGACY.items():
            st = SKELETON_DEFAULT_STYLE[sk]
            self.assertEqual(st.source, "skeleton_default")
            self.assertEqual(st.name_size, want["name"])
            self.assertEqual(st.title_size, want["title"])
            self.assertEqual(st.org_size, want["org"])
            self.assertEqual(st.role_size, want["role"])
            self.assertEqual(st.body_size, want["body"])
            self.assertEqual(st.meta_size, want["meta"])
            self.assertEqual(st.font_family, "微软雅黑")

    def test_02_banner_legacy_sizes_and_spacing(self):
        d = build_document(fake_blocks(), "banner_card")
        want = LEGACY["banner_card"]
        self.assertEqual(size_of_para(d, "张三")[0], want["name"])
        self.assertEqual(size_of_para(d, "测试岗位")[0], want["intent"])
        self.assertEqual(size_of_para(d, "zhangsan@example.com")[0],
                         want["contact"])
        self.assertEqual(size_of_para(d, "教育经历")[0], want["title"])
        hp = find_para(d, "某某有限公司")
        sizes = dict(run_sizes(hp))
        self.assertAlmostEqual(sizes["某某有限公司"], want["org"], delta=0.005)
        self.assertAlmostEqual(sizes["  ·  测试实习生"], want["role"], delta=0.005)
        self.assertAlmostEqual(sizes["\t2025.06 - 2025.09"], want["meta"],
                               delta=0.005)
        # 行距：经历表头 1.05；正文 1.08；板块标题 1.0
        self.assertAlmostEqual(hp.paragraph_format.line_spacing,
                               want["line_exp"], delta=0.005)
        self.assertAlmostEqual(
            find_para(d, "虚构的个人优势").paragraph_format.line_spacing,
            want["line_body"], delta=0.005)
        # 字体三件套全部落到微软雅黑
        self.assertEqual(first_run_font(hp), ("微软雅黑", "微软雅黑"))

    def test_03_minimal_legacy_sizes(self):
        d = build_document(fake_blocks(), "single_column_minimal")
        want = LEGACY["single_column_minimal"]
        self.assertEqual(size_of_para(d, "张三")[0], want["name"])
        self.assertEqual(size_of_para(d, "测试岗位")[0], want["intent"])
        self.assertAlmostEqual(
            find_para(d, "虚构的个人优势").paragraph_format.line_spacing,
            want["line_body"], delta=0.005)
        # minimal bullet 行距为独立的 1.08
        bp = find_para(d, "完成了一项虚构的测试工作")
        self.assertAlmostEqual(bp.paragraph_format.line_spacing, 1.08,
                               delta=0.005)

    def test_04_sidebar_legacy_sizes_and_spacing(self):
        d = build_document(fake_blocks(), "two_column_sidebar")
        want = LEGACY["two_column_sidebar"]
        self.assertEqual(size_of_para(d, "张三")[0], want["name"])
        self.assertEqual(size_of_para(d, "测试岗位")[0], want["intent"])
        self.assertEqual(size_of_para(d, "技能甲")[0], want["side_body"])
        # 右栏正文 1.08；cell 内 bullet 1.06；侧栏短行 1.05
        self.assertAlmostEqual(
            find_para(d, "虚构的个人优势").paragraph_format.line_spacing,
            want["line_body"], delta=0.005)
        self.assertAlmostEqual(
            find_para(d, "· 完成了一项虚构的测试工作")
                .paragraph_format.line_spacing,
            want["line_bullet"], delta=0.005)
        self.assertAlmostEqual(
            find_para(d, "张三").paragraph_format.line_spacing,
            want["line_side"], delta=0.005)

    def test_05_no_numeric_size_literals_in_skeleton_calls(self):
        """源码守卫：add_text(...) 不允许再传数字字号（视觉决策必须来自 style）。"""
        src = (ROOT / "src" / "resume_generator" / "skeletons.py").read_text(
            encoding="utf-8")
        # 去掉注释行，避免误报
        code = "\n".join(
            ln for ln in src.splitlines() if not ln.strip().startswith("#"))
        offenders = [
            ln.strip() for ln in code.splitlines()
            if re.search(r"add_text\([^)]*,\s*\d+(?:\.\d+)?\s*[,)]", ln)
        ]
        self.assertEqual(offenders, [],
                         f"发现数字字号字面量：{offenders}")


def half_pt(v):
    """python-docx 把 pt 截断到 w:sz 半磅整数网格后的回读值。"""
    import math
    return math.floor(v * 2) / 2


class TestTypographySpecDriven(unittest.TestCase):
    """B. Spec 路径：取值来自 TypographySpec，缺口显式回退。"""

    def test_06_font_token_mapping(self):
        self.assertEqual(resolve_font_family("sans_zh"), "微软雅黑")
        self.assertEqual(resolve_font_family("Microsoft YaHei"), "微软雅黑")
        self.assertEqual(resolve_font_family(None), "微软雅黑")
        self.assertEqual(resolve_font_family("SomeOther Font"),
                         "SomeOther Font")

    def test_07_p1_spec_typography_applied(self):
        from resume_generator.design import build_p1_spec
        d = build_document(fake_blocks(), design_spec=build_p1_spec())

        # 样本A实测字号（按半磅网格截断后比对）
        self.assertAlmostEqual(size_of_para(d, "张三")[0],
                               half_pt(19.56), delta=0.005)
        hp = find_para(d, "某某有限公司")
        sizes = dict(run_sizes(hp))
        self.assertAlmostEqual(sizes["某某有限公司"], half_pt(10.56),
                               delta=0.005)
        # Phase 3B-4：角色分隔符已 Spec 驱动（experience.role.separator="｜"），
        # 不再使用 banner Profile 的 "  ·  "；字号断言不变。
        self.assertAlmostEqual(sizes["｜测试实习生"], half_pt(9.96),
                               delta=0.005)
        self.assertAlmostEqual(sizes["\t2025.06 - 2025.09"], half_pt(9.0),
                               delta=0.005)
        self.assertAlmostEqual(size_of_para(d, "教育经历")[0],
                               half_pt(11.52), delta=0.005)
        # 深块联系方式取 header_meta_sizes[0] = 8.04
        self.assertAlmostEqual(
            size_of_para(d, "zhangsan@example.com")[0],
            half_pt(8.04), delta=0.005)
        # 正文 9.48 + 1.9 倍宽行距；标题行 1.05
        sp = find_para(d, "虚构的个人优势")
        self.assertAlmostEqual(run_sizes(sp)[0][1], half_pt(9.48),
                               delta=0.005)
        self.assertAlmostEqual(sp.paragraph_format.line_spacing, 1.9,
                               delta=0.005)
        self.assertAlmostEqual(
            find_para(d, "教育经历").paragraph_format.line_spacing,
            1.05, delta=0.005)
        self.assertAlmostEqual(hp.paragraph_format.line_spacing, 1.05,
                               delta=0.005)
        self.assertEqual(first_run_font(hp), ("微软雅黑", "微软雅黑"))

    def test_08_p2_spec_typography_and_fallbacks(self):
        from resume_generator.design import build_p2_spec
        spec = build_p2_spec()
        st = resolve_style_for_spec(
            spec,
            SKELETON_DEFAULT_STYLE["two_column_sidebar"],
            "two_column_sidebar",
            resolve_palette(None, "two_column_sidebar"))
        self.assertEqual(st.source, "design_spec")

        d = build_document(fake_blocks(), design_spec=spec)
        self.assertAlmostEqual(size_of_para(d, "张三")[0], half_pt(26.0),
                               delta=0.005)
        self.assertAlmostEqual(size_of_para(d, "教育经历")[0], half_pt(14.0),
                               delta=0.005)
        hp = find_para(d, "某某有限公司")
        sizes = dict(run_sizes(hp))
        self.assertAlmostEqual(sizes["  ·  测试实习生"], half_pt(10.0),
                               delta=0.005)
        self.assertAlmostEqual(
            find_para(d, "虚构的个人优势").paragraph_format.line_spacing,
            1.5, delta=0.005)
        # 侧栏短行行距 1.4
        self.assertAlmostEqual(
            find_para(d, "技能甲").paragraph_format.line_spacing,
            1.4, delta=0.005)
        # 侧栏正文 9pt
        self.assertAlmostEqual(size_of_para(d, "技能甲")[0], half_pt(9.0),
                               delta=0.005)

        # 缺口回退（不编造）：
        # weights.role="undetermined" → 骨架粗体兜底
        self.assertTrue(st.w_role)
        # sidebar_group_label_size undetermined → 9.0
        self.assertEqual(st.side_label_size, 9.0)
        # intent_size 在 TypographySpec 无字段 → 8.5（sidebar 原值）
        self.assertEqual(st.intent_size, 8.5)
        # header_meta_sizes=None → 侧栏联系方式沿用 8.5
        self.assertEqual(st.contact_size, 8.5)
        # 回退登记可审计
        self.assertIn("typography.weights.role", st.spec_fallbacks)
        self.assertTrue(
            any("intent_size" in k for k in st.spec_fallbacks))
        self.assertIn("typography.sidebar_group_label_size",
                      st.spec_fallbacks)

    def test_09_p2_role_undetermined_renders_bold(self):
        """缺口兜底的字重必须真实落到 run（w:b），而非只存在参数包里。"""
        from docx.oxml.ns import qn as _qn
        from resume_generator.design import build_p2_spec
        d = build_document(fake_blocks(), design_spec=build_p2_spec())
        hp = find_para(d, "某某有限公司")
        role_run = next(r for r in hp.runs if "测试实习生" in r.text)
        self.assertIsNotNone(role_run._r.find(_qn("w:rPr")).find(_qn("w:b")))


if __name__ == "__main__":
    unittest.main(verbosity=2)

